"""資料作成・変換表 Phase 4: アドバリュー 依頼表 → 報告書化＋京阪南版エリア被り判定。

入力2つ:
- アドバリュー依頼表（A=GeoCode / B=市区名 / C=町丁目名(大久保町1) / D=世帯数15 /
  E=共同住宅主世帯数15 / F=暫定部数 / G=クライアント①○× / H=クライアント②○× /
  I=配布日メモ / J=併配チラシ暫定部数 / K=備考(店舗名) / L=投禁物件 / M=配布日 / N=重複）
- 京阪南版 配送管理表（京阪北と同じ17列。N=町界名 / O=街区(番地)名称 / I=担当地区6桁）

やること: 依頼表の各「町丁目名」を 町名＋丁目 に分解し、京阪南の 町界名＋街区 に照合して
担当地区を割り出す（マッチ=京阪南と被り／未マッチ=被らない）。丁目パーサは現行マクロの
取りこぼしバグ（「1丁目・2丁目11～34」等で2丁目以降を落とす）を直して正確に判定する。
出力=依頼表の列＋「担当地区(被り)」＋「被り/非被り」の報告書Excel。
"""
import datetime as _dt
import io
import re

import openpyxl

from common.atehagi import read_uploaded
from common.excel_io import freeze_xlsx_bytes

# ===== 週次の案件区分（依頼⑥・2026-08-27 大橋様ご指摘） =====
# アドバリューは週ごとに依頼が来て「8-1」「8-2」…と区分して管理している。
# 登録時に必ず週を選ばせ、9月になれば自動で「9-1」から始まるようにしたい。
#
# 🔴 設計判断: 画面で選ぶ選択肢は月に依存しない「1週目〜5週目」にし、
# 月の部分は **登録する伝票自身の日付** から計算して "8-1" の形で保存する。
#   ・選択肢に月を焼き込むと、Streamlit のフォームは日付を変えても再実行
#     されないため「10月の伝票を入力しているのに選択肢は9月のまま」になる。
#   ・保存する値の形は従来の手入力("8-1")と同じなので、8/20以前に登録した
#     データとそのまま同じグループに並ぶ(移行作業が要らない)。
#   ・月が変われば伝票の日付が変わるので、選択肢を作り直さなくても
#     9-1 → 10-1 と自動で繰り上がる。
WEEK_CHOICES = ("1週目", "2週目", "3週目", "4週目", "5週目")
WEEK_PLACEHOLDER = "（選択してください）"
_WEEK_LABEL_RE = re.compile(r"^\s*(\d{1,2})\s*-\s*([1-9])\s*$")


def week_label(month, week) -> str:
    """月と週から案件区分の文字列を作る。8月の1週目→"8-1"。"""
    return f"{int(month)}-{int(week)}"


def week_index(choice):
    """「3週目」→3。未選択・プレースホルダ・読めない値は None。"""
    m = re.match(r"^\s*([1-9])\s*週目\s*$", str(choice or ""))
    return int(m.group(1)) if m else None


def month_of(value, *, today=None) -> int:
    """伝票の日付(YYYY-MM-DD / YYYY-MM / date)から月を取り出す。

    読めないときは登録日(today)の月に倒す。ここで例外を投げると
    「日付を空にしたまま登録しようとしたら画面が落ちた」になるため。
    """
    if isinstance(value, (_dt.date, _dt.datetime)):
        return value.month
    m = re.match(r"^\s*(\d{4})\D(\d{1,2})", str(value or ""))
    if m:
        month = int(m.group(2))
        if 1 <= month <= 12:
            return month
    if today is None:
        return _dt.date.today().month
    if isinstance(today, (_dt.date, _dt.datetime)):
        return today.month
    return int(str(today)[5:7])


def week_label_for(date_value, choice, *, today=None):
    """伝票の日付と選んだ週から案件区分を作る。週が未選択なら None。"""
    week = week_index(choice)
    if week is None:
        return None
    return week_label(month_of(date_value, today=today), week)


def parse_week_label(label):
    """"8-1" → (8, 1)。週の区分でない文字列(手入力の案件名など)は None。"""
    m = _WEEK_LABEL_RE.match(str(label or ""))
    if not m:
        return None
    month = int(m.group(1))
    return (month, int(m.group(2))) if 1 <= month <= 12 else None


def sort_week_labels(labels):
    """案件区分を 月→週 の順に並べる。

    🔴 素の文字列順だと "10-1" が "8-1" より前に来て、月をまたいだ瞬間に
    並びが壊れる。週として読めないラベル(手入力の案件名)は末尾に五十音順で置く
    (捨てない＝どこにも出ない区分を作らない)。
    """
    weeks = sorted((l for l in labels if parse_week_label(l)),
                   key=lambda l: parse_week_label(l))
    others = sorted(l for l in labels if not parse_week_label(l))
    return weeks + others


ADVALUE_PROJECT = "アドバリュー"
SONOTA_PROJECT = "その他"
# 登録画面で案件区分(other_label)を使う案件。ここに無い案件では区分を保存しない。
CASE_LABEL_PROJECTS = (ADVALUE_PROJECT, SONOTA_PROJECT)
WEEK_REQUIRED_MESSAGE = "アドバリューは週（1週目〜5週目）を選んでください。"


def resolve_case_label(project_name, *, date_value, week_choice, free_text, today=None):
    """登録する行の案件区分(other_label)を決めて (値, エラー文) で返す。

    ・アドバリュー … 週の選択が必須。伝票の日付の月と組み合わせて "8-1" にする。
      自由入力は使わない(「8-1」「8月1週」「8_1」と表記が割れるのを防ぐため)。
    ・その他 … これまでどおり自由入力(何の案件か)をそのまま。空でもよい。
    ・それ以外の案件 … 区分は持たない(None)。
    エラー文が返ったときは登録してはいけない。
    """
    name = str(project_name or "").strip()
    if name == ADVALUE_PROJECT:
        label = week_label_for(date_value, week_choice, today=today)
        if label is None:
            return None, WEEK_REQUIRED_MESSAGE
        return label, None
    if name == SONOTA_PROJECT:
        return (str(free_text or "").strip() or None), None
    return None, None


def week_display(label) -> str:
    """画面に出すときの表記。"8-1" → "8月 1週目"。週でなければそのまま。"""
    parsed = parse_week_label(label)
    if not parsed:
        return str(label or "")
    month, week = parsed
    return f"{month}月 {week}週目"

_ZEN2HAN = str.maketrans("０１２３４５６７８９", "0123456789")


def _z2h(s) -> str:
    return "" if s is None else str(s).translate(_ZEN2HAN)


def normalize_town(s) -> str:
    """町名の突合キー: 全角数字→半角、空白除去。"""
    return re.sub(r"\s|　", "", _z2h(s)).strip()


def parse_choume_from_gaiku(gaiku):
    """京阪南マスタの「街区（番地）名称」から、含まれる丁目番号の集合を返す。

    例: '1丁目1～21'→{1} / '2・3丁目'→{2,3} / '4・5丁目'→{4,5} /
        '6丁目51～57・7丁目'→{6,7} / '1丁目・2丁目11～34'→{1,2}
    丁目が無い（空/全域）→ 空集合（＝その町の全丁目を覆うワイルドカード扱い）。
    """
    s = _z2h(gaiku)
    if s.strip() == "":
        return set()
    s = re.sub(r"[、，]", "・", s)               # 区切りを ・ に統一
    s = re.sub(r"[～〜~－ー–—-]", "~", s)         # 範囲記号を ~ に統一
    parts = s.split("丁目")
    if len(parts) == 1:
        return set()                             # 丁目表記なし＝ワイルドカード
    choume = set()
    k = len(parts) - 1                           # 丁目 の出現回数
    for i in range(k):
        seg = parts[i]
        tokens = [t.strip() for t in seg.split("・")]
        if i == 0:
            # 先頭: 丁目の直前は全部が丁目番号（番地は丁目より後にしか来ない）
            for t in tokens:
                if re.fullmatch(r"\d+", t):
                    choume.add(int(t))
        else:
            # 途中: 「前の丁目の番地レンジ ・ 次の丁目番号」→ 末尾の素の数字だけが次の丁目
            for t in reversed(tokens):
                if re.fullmatch(r"\d+", t):
                    choume.add(int(t))
                else:
                    break
    return choume


def split_town_choume(choume_name):
    """依頼表の町丁目名を (町名, 丁目番号 or None) に分解する。例 '大久保町1'→('大久保町',1)。"""
    s = _z2h(choume_name).strip()
    s = re.sub(r"丁目$", "", s)
    m = re.match(r"^(.*?)(\d+)$", s)
    if m and m.group(1):
        return (m.group(1), int(m.group(2)))
    return (s, None)


_MASTER_HEADER_KEY = "担当地区"
_COL_CHIKU = 9      # I 担当地区
_COL_CHOUKAI = 14   # N 町界名
_COL_GAIKU = 15     # O 街区（番地）名称


def build_minami_index(minami_table):
    """京阪南配送管理表から {normalize(町名): [(丁目集合, 担当地区コード), ...]} を作る。"""
    hidx = None
    for i, row in enumerate(minami_table):
        if row and any(str(c).strip() == _MASTER_HEADER_KEY for c in row if c is not None):
            hidx = i
            break
    if hidx is None:
        raise ValueError("京阪南 配送管理表のヘッダー行（担当地区）が見つかりません")

    def cell(row, col1):
        c = col1 - 1
        return row[c] if row and c < len(row) else None

    index = {}
    seen = set()   # (町名, 街区, 担当地区) の重複行(=同地区の複数チラシ行)を1つに
    for row in minami_table[hidx + 1:]:
        if row is None:
            continue
        chiku = cell(row, _COL_CHIKU)
        choukai = cell(row, _COL_CHOUKAI)
        if chiku is None or str(chiku).strip() == "" or choukai is None:
            continue
        code = str(chiku).strip()
        town = str(choukai).strip()
        gaiku = cell(row, _COL_GAIKU)
        dedup = (town, str(gaiku), code)
        if dedup in seen:
            continue
        seen.add(dedup)
        index.setdefault(normalize_town(town), []).append(
            (parse_choume_from_gaiku(gaiku), code))
    return index


def match_area(town, choume, index):
    """町名＋丁目 に一致する京阪南の担当地区コード（複数可・重複除去・順序保持）を返す。

    町名は完全一致を優先し、無ければ「町」サフィックスの有無を吸収して照合する
    （依頼表「大日」＝京阪南「大日町」のような省略に対応）。
    """
    key = normalize_town(town)
    hits = index.get(key)
    if hits is None and not key.endswith("町"):
        hits = index.get(key + "町")            # 「大日」→「大日町」
    if hits is None and key.endswith("町"):
        hits = index.get(key[:-1])              # 逆方向も一応
    hits = hits or []
    codes = []
    for choume_set, code in hits:
        if choume is None or not choume_set or choume in choume_set:
            if code not in codes:
                codes.append(code)
    return codes


_IRAI_HEADER_KEY = "町丁目名"
_IRAI_COLS = {
    "geocode": 1, "shiku": 2, "choume_name": 3, "setai": 4, "kyodo_setai": 5,
    "zantei": 6, "client1": 7, "client2": 8, "haifubi_memo": 9, "heihai": 10,
    "biko": 11, "toukin": 12, "haifubi": 13,
}


def parse_advalue_irai(table):
    """アドバリュー依頼表を行dictのリストに正規化する。"""
    hidx = None
    for i, row in enumerate(table):
        if row and any(str(c).strip() == _IRAI_HEADER_KEY for c in row if c is not None):
            hidx = i
            break
    if hidx is None:
        raise ValueError("依頼表のヘッダー行（町丁目名）が見つかりません")

    def cell(row, col1):
        c = col1 - 1
        return row[c] if row and c < len(row) else None

    def s(v):
        return "" if v is None else str(v).strip()

    rows = []
    for row in table[hidx + 1:]:
        if row is None:
            continue
        name = s(cell(row, _IRAI_COLS["choume_name"]))
        if name == "" or name in ("合計", "合 計"):
            break
        rows.append({k: cell(row, col1) for k, col1 in _IRAI_COLS.items()})
    return rows


_REPORT_HEADERS = [
    "町丁目名", "市区名", "担当地区(被り)", "被り", "世帯数15", "暫定部数",
    "クライアント①", "クライアント②", "配布日メモ", "併配部数", "店舗名", "投禁物件", "配布日",
]


def build_report_rows(irai_rows, minami_index):
    """依頼表＋京阪南突合の結果を、報告書の行(dictリスト)にする。"""
    out = []
    for r in irai_rows:
        name = "" if r["choume_name"] is None else str(r["choume_name"]).strip()
        town, choume = split_town_choume(name)
        codes = match_area(town, choume, minami_index)
        out.append({
            "町丁目名": name,
            "市区名": r["shiku"],
            "担当地区(被り)": "、".join(codes) if codes else "未マッチ",
            "被り": "被り" if codes else "非被り",
            "世帯数15": r["setai"],
            "暫定部数": r["zantei"],
            "クライアント①": r["client1"],
            "クライアント②": r["client2"],
            "配布日メモ": r["haifubi_memo"],
            "併配部数": r["heihai"],
            "店舗名": r["biko"],
            "投禁物件": r["toukin"],
            "配布日": r["haifubi"],
        })
    return out


def build_advalue_report_workbook(report_rows, *, case_label=None) -> bytes:
    """報告書Excel（1シート＝ヘッダー＋明細）を bytes で返す。

    アドバリューは週ごとに依頼が来て『8-1』『8-2』『8-3』のように案件を
    分けて管理している(2026-08-19 大橋様ご指摘)。case_label を渡すと、
    どの案件の報告書か分かるよう先頭に見出し行を入れる。渡さなければ
    これまでどおりヘッダーがA1から始まる(後方互換)。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "アドバリュー報告書"
    header_row = 1
    if case_label:
        ws.cell(row=1, column=1, value=f"アドバリュー エリア被り報告書　{case_label}")
        ws.cell(row=1, column=1).font = openpyxl.styles.Font(bold=True, size=14)
        header_row = 2
    for col, h in enumerate(_REPORT_HEADERS, start=1):
        ws.cell(row=header_row, column=col, value=h)
    for i, r in enumerate(report_rows):
        for col, h in enumerate(_REPORT_HEADERS, start=1):
            ws.cell(row=header_row + 1 + i, column=col, value=r.get(h))
    for col in range(1, len(_REPORT_HEADERS) + 1):
        ws.cell(row=header_row, column=col).font = openpyxl.styles.Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return freeze_xlsx_bytes(buf.getvalue())


_FILENAME_UNSAFE = re.compile(r'[\\/:*?"<>|]')


def advalue_filename(case_label=None) -> str:
    """報告書のファイル名。案件名を渡すと、週ごとに生成しても上書きせず
    分けて保存できるようファイル名にも入れる(2026-08-19 大橋様ご指摘)。"""
    base = "アドバリュー_エリア被り報告書"
    if case_label:
        safe = _FILENAME_UNSAFE.sub("_", str(case_label).strip())
        if safe:
            base += f"_{safe}"
    return base + ".xlsx"


def overlap_summary(report_rows):
    """被り/非被りの件数サマリ。"""
    hit = sum(1 for r in report_rows if r["被り"] == "被り")
    return {"total": len(report_rows), "overlap": hit, "non_overlap": len(report_rows) - hit}


def read_irai(name, data):
    return parse_advalue_irai(read_uploaded(name, data))


def read_minami_index(name, data):
    return build_minami_index(read_uploaded(name, data))
