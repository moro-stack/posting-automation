import re
import zipfile
from io import BytesIO

import pandas as pd


# ===== xlsx の出力を「同じ内容なら同じバイト列」に揃える =====
# xlsx は ZIP で、openpyxl/pandas は保存のたびに「今の時刻」を
#   ・docProps/core.xml の作成/更新日時
#   ・ZIP内の各エントリのタイムスタンプ
# に書き込む。そのため同じ内容でも保存のたびにバイト列が変わる。
# Streamlit のダウンロードURL(/media/<hash>.xlsx)はファイル内容のハッシュで決まるため、
# バイト列が毎回変わると再実行のたびに別URLが発行され、ブラウザが取得しようとしている
# 古いURLがサーバー側から破棄されて404になる(＝「Excelがダウンロードされない」)。
# 時刻を固定して内容を安定させ、URLが変わらないようにする。
_ZIP_DATE = (1980, 1, 1, 0, 0, 0)   # ZIPが表現できる最小の日時
_CORE_XML = "docProps/core.xml"
_DCTERMS_RE = re.compile(
    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)")
_FIXED_STAMP = b"2020-01-01T00:00:00Z"


def freeze_xlsx_bytes(data: bytes) -> bytes:
    """xlsx(ZIP)内のタイムスタンプと更新日時を固定値に揃えて詰め直す。
    セルの値・書式は一切変えない。同じ内容なら毎回同じバイト列になる。"""
    src = zipfile.ZipFile(BytesIO(data))
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == _CORE_XML:
                content = _DCTERMS_RE.sub(rb"\g<1>" + _FIXED_STAMP + rb"\g<2>", content)
            frozen = zipfile.ZipInfo(info.filename, date_time=_ZIP_DATE)
            frozen.compress_type = zipfile.ZIP_DEFLATED
            frozen.external_attr = info.external_attr
            frozen.internal_attr = info.internal_attr
            frozen.create_system = info.create_system
            dst.writestr(frozen, content)
    return out.getvalue()


# テンプレに同梱されている円グラフ一式(描画+グラフ本体+リレーション)
_CHART_PARTS = (
    "xl/drawings/drawing1.xml",
    "xl/drawings/_rels/drawing1.xml.rels",
    "xl/charts/chart1.xml",
    "xl/charts/chart2.xml",
)
_CT_OVERRIDES = (
    '<Override PartName="/xl/drawings/drawing1.xml"'
    ' ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
    '<Override PartName="/xl/charts/chart1.xml"'
    ' ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>'
    '<Override PartName="/xl/charts/chart2.xml"'
    ' ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>'
)
_DRAW_REL = (
    '<Relationship Id="rIdDraw1"'
    ' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"'
    ' Target="/xl/drawings/drawing1.xml"/>'
)
_DRAW_REF = (
    '<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    ' r:id="rIdDraw1"/>'
)


def inject_template_charts(saved_bytes: bytes, template_path: str, new_sheet_title: str,
                           report_sheet_part: str = "xl/worksheets/sheet1.xml") -> bytes:
    """openpyxl保存済みxlsxに、テンプレ同梱の円グラフ一式をそのまま注入する。

    openpyxlはテンプレートの既存グラフを再保存すると体裁が変わったりExcelが破損判定する。
    そこで保存時はグラフを外しておき(report側で sheet._charts=[])、最後にテンプレの
    グラフXMLをバイト等価で注入する。シート名参照(<f>のシート名)だけを新シート名へ置換
    するので、見た目(大きさ・淵・フォント・色)はテンプレと完全一致する。

    Args:
        saved_bytes: グラフを外した状態でopenpyxlが保存したxlsxのバイト列。
        template_path: 円グラフを持つテンプレート.xlsxのパス。
        new_sheet_title: 報告書シートの現在のタブ名(<f>参照をこの名前に書き換える)。
        report_sheet_part: 報告書シートのworksheet XMLパス(通常 sheet1.xml)。
    """
    with zipfile.ZipFile(template_path) as tmpl:
        chart1 = tmpl.read("xl/charts/chart1.xml").decode("utf-8")
        m = re.search(r"<c:f>([^!<]+)!", chart1) or re.search(r"<f>([^!<]+)!", chart1)
        old_sheet = m.group(1) if m else None
        parts = {}
        for name in _CHART_PARTS:
            data = tmpl.read(name)
            if name.startswith("xl/charts/chart") and old_sheet:
                # シート名参照だけを新シート名へ(見た目XMLには触れない)
                data = data.decode("utf-8").replace(old_sheet, new_sheet_title).encode("utf-8")
            parts[name] = data

    src = zipfile.ZipFile(BytesIO(saved_bytes))
    existing = set(src.namelist())
    rels_part = report_sheet_part.replace("xl/worksheets/", "xl/worksheets/_rels/") + ".rels"
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.decode("utf-8").replace("</Types>", _CT_OVERRIDES + "</Types>").encode("utf-8")
            elif item.filename == report_sheet_part and "<drawing " not in data.decode("utf-8", "replace"):
                text = data.decode("utf-8")
                # <drawing> はワークシート末尾付近。extLstがあればその前に入れる。
                anchor = "<extLst>" if "<extLst>" in text else "</worksheet>"
                data = text.replace(anchor, _DRAW_REF + anchor, 1).encode("utf-8")
            elif item.filename == rels_part:
                data = data.decode("utf-8").replace("</Relationships>", _DRAW_REL + "</Relationships>").encode("utf-8")
            zout.writestr(item, data)
        if rels_part not in existing:
            zout.writestr(
                rels_part,
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                + _DRAW_REL + "</Relationships>",
            )
        for name, data in parts.items():
            zout.writestr(name, data)
    return out.getvalue()


# Excel がシート名に受け付けない文字。
# 🔴 半角だけでなく全角も同じ扱いで弾かれる。Excel は全角を半角に正規化してから
#    判定するため、禁止文字を「全角の双子に寄せて回避する」ことはできない。
#    2026-08-10、仕分け表で ／(U+FF0F) を使っていたシートが Excel に修復され
#    『回復済み_Sheet1』に化けた（修復メッセージ＝/xl/workbook.xml のワークシートのプロパティ）。
#    シート名を作る処理は、必ずこの定数を使って判定すること。
EXCEL_NG_SHEET_CHARS = "[]:*?/" + "\\" + "［］：＊？／＼"


def sanitize_sheet_chars(title, replacement="_") -> str:
    """Excelがシート名に受け付けない文字を replacement に置き換える（半角・全角とも）。"""
    return "".join(replacement if ch in EXCEL_NG_SHEET_CHARS else ch
                   for ch in str(title))


def safe_sheet_title(title: str, existing=()) -> str:
    """Excelのシート名禁止文字を除き31文字に丸め、既存名と重複しないようにする。"""
    cleaned = sanitize_sheet_chars(title, "").strip() or "Sheet"
    cleaned = cleaned[:31]
    if cleaned not in existing:
        return cleaned
    # 重複時は連番を付与(31文字以内を維持)
    for i in range(2, 100):
        suffix = f"_{i}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        if candidate not in existing:
            return candidate
    return cleaned


# 旧名。シート名を作る処理はこの1つに集約する方針(2026-08-10 f00515c)なので、
# 既存の呼び出し・テストが動くよう別名を残す。
_safe_sheet_title = safe_sheet_title


def append_dataframe_sheet(workbook, title: str, df: "pd.DataFrame", row_height=None):
    """DataFrame を新しいシートとして workbook に追加する(1行目=ヘッダー)。

    返却リスト等、アップロードした表をそのまま別シートに同梱するのに使う。
    NaN は空欄に、numpy型はPython標準型に変換して書き込む。文字列は電話番号等の
    先頭0が消えないよう、明示的にテキスト書式(@)で書き込む。row_height 指定で行高を統一。
    """
    from openpyxl.utils import get_column_letter
    safe = safe_sheet_title(title, existing=set(workbook.sheetnames))
    ws = workbook.create_sheet(title=safe)
    ws.append([str(c) for c in df.columns])
    for row in df.itertuples(index=False, name=None):
        out = []
        for v in row:
            if v is None or (isinstance(v, float) and pd.isna(v)):
                out.append("")
            elif hasattr(v, "item"):  # numpy スカラ → Python標準型
                out.append(v.item())
            else:
                out.append(v)
        ws.append(out)
    # 文字列セルはテキスト書式(@)にして、先頭0付き電話番号などが数値化・0落ちしないようにする
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, str):
                cell.number_format = "@"
    if row_height is not None:
        for r in range(1, ws.max_row + 1):
            ws.row_dimensions[r].height = row_height
    return ws


def write_status_counts_to_sheet(sheet, status_counts: dict, start_row: int, start_col: int, end_col: int, max_rows: int) -> None:
    """Clear a template range and write status-count pairs into the first two columns.

    This function clears columns [start_col, end_col] over [start_row, start_row + max_rows)
    to remove any stale data or decorative elements from the template area. However, data is
    always written to exactly two columns: start_col (status name) and start_col + 1 (count),
    regardless of end_col. This design matches the business requirement of clearing a wider
    "架電結果表エリア" template while only re-writing two logical data columns.

    Entries in status_counts beyond max_rows are silently omitted (not written and no exception
    raised). The caller is responsible for ensuring max_rows is large enough for the
    expected template, as truncation is by design and avoids raising/crashing on oversized input.

    Args:
        sheet: openpyxl Worksheet object to write to.
        status_counts: Dictionary mapping status name (str) to count (int or numeric).
        start_row: First row index (1-indexed) where the template area begins.
        start_col: First column index (1-indexed) where to write status names.
        end_col: Last column index (1-indexed) to clear in the template area.
                 Only defines the CLEAR range, not the WRITE range.
        max_rows: Maximum rows to write; any entries beyond this are truncated silently.

    Returns:
        None. Modifies sheet in-place.
    """
    for row in range(start_row, start_row + max_rows):
        for col in range(start_col, end_col + 1):
            sheet.cell(row=row, column=col).value = None

    for offset, (status, count) in enumerate(status_counts.items()):
        row = start_row + offset
        if row >= start_row + max_rows:
            break
        sheet.cell(row=row, column=start_col).value = status
        sheet.cell(row=row, column=start_col + 1).value = count


def workbook_to_bytes(workbook) -> bytes:
    """Serialize an in-memory openpyxl Workbook to bytes.

    Converts a Workbook object (with all its sheets and formatting) into a
    binary Excel file format (.xlsx) suitable for file downloads or transmission.
    Useful with Streamlit's download_button or similar web download utilities.

    Args:
        workbook: openpyxl Workbook object.

    Returns:
        bytes: Binary Excel file content (.xlsx format).
    """
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
