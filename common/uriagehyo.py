"""アプリのデータを、大阪支社の「会議で使う売上表」(手作りの月次台帳)と
同じ列並びで書き出す(依頼⑦・2026-08-19 大橋様ご指摘)。

会議用売上表(KPS大阪売上表.xlsx)は「ぱど」「チラシ」「仕分け」の3区分の
部数・売上まで持つが、この区分はアプリのどこにも記録されていない。
アプリが持っている値(売上合計・配布原価・交通費(駐車場代)・飲み物代)だけを
同じ列に自動で入れ、区分の無い列と備考は空欄のまま出す。

2026-08-20 オーナー判断＝実ファイル(19シートの財務台帳)への直接書き込みは
せず、都度エクスポートしたものをコピー＆ペーストしてもらう運用にする。
"""
import io

import openpyxl

from common import excel_io

HEADERS = [
    "発行号", "版名", "ぱど部数", "ぱど売上（税抜）", "チラシ部数", "チラシ売上（税抜）",
    "仕分け部数", "仕分け売上（税抜）", "その他", "売上合計（税抜）", "売上合計（税込）",
    "備考", "交通費（駐車場代含）", "飲み物", "配布原価（税込）", "原価合計（税込）",
]


def build_row(*, hakko_gou, ban_mei, uriage_zeikomi, genka_goukei_zeikomi,
             koutsuhi=0, nomimono=0):
    """1号ぶんのデータを売上表の1行(dict)にする。

    genka_goukei_zeikomi は号別明細の「配布原価(税込)」＝アプリの原価集計の
    総額(小口の交通費・飲み物代も含む)をそのまま渡す。実物の売上表は
    「配布原価」列を交通費・飲み物を除いた額、「原価合計」列をその総額に
    分けているため、ここで差し引いて2列に振り分ける
    (配布原価＝総額－交通費－飲み物、原価合計＝総額のまま)。

    売上合計(税抜)は税込から逆算(10%)。
    ぱど/チラシ/仕分け/その他/備考はアプリに区分の記録が無いため空欄で返す。
    交通費/飲み物は実物の表にならい、0円(記録が無い)なら空欄にする。
    """
    uriage_zeikomi = int(uriage_zeikomi or 0)
    genka_goukei = int(genka_goukei_zeikomi or 0)
    koutsuhi = int(koutsuhi or 0)
    nomimono = int(nomimono or 0)
    uriage_zeinuki = round(uriage_zeikomi / 1.1)
    haifu_genka = genka_goukei - koutsuhi - nomimono
    return {
        "発行号": hakko_gou, "版名": ban_mei,
        "ぱど部数": None, "ぱど売上（税抜）": None,
        "チラシ部数": None, "チラシ売上（税抜）": None,
        "仕分け部数": None, "仕分け売上（税抜）": None,
        "その他": None,
        "売上合計（税抜）": uriage_zeinuki, "売上合計（税込）": uriage_zeikomi,
        "備考": None,
        "交通費（駐車場代含）": koutsuhi or None, "飲み物": nomimono or None,
        "配布原価（税込）": haifu_genka, "原価合計（税込）": genka_goukei,
    }


# ===== 会議用売上表のセクション（依頼⑦・2026-08-27 大橋様） =====
# 新テンプレート(KPS(大阪)原価売上表テンプレート.xlsx)は
#   京阪南／京阪北 → 単独チラシ → アド・バリュー(週ごと) → リビング → 集計
# というセクション構成になっている。案件マスタの名前をこのセクションに割り当てる。
SECTION_PADO_MINAMI = "京阪南"
SECTION_PADO_KITA = "京阪北"
SECTION_TANDOKU = "単独"        # 単独チラシ＝案件「その他」(件数ぶん行を増やす)
SECTION_ADVALUE = "アドバリュー"
SECTION_LIVING = "リビング"
SECTION_UNKNOWN = "未分類"      # どれにも当てはまらない案件・案件未設定

_PROJECT_SECTIONS = {
    "関西ぱど：京阪南版": SECTION_PADO_MINAMI,
    "関西ぱど：京阪北版": SECTION_PADO_KITA,
    "その他": SECTION_TANDOKU,
    "アドバリュー": SECTION_ADVALUE,
    "リビングプロシード": SECTION_LIVING,
}


def section_of(label) -> str:
    """売上表の見出しラベルから、どのセクションの行かを決める。

    ラベルは build_bulk_rows が作る「案件名」または「案件名　区分」。
    🔴 知らない案件名は捨てずに SECTION_UNKNOWN に入れる。案件マスタが増えた
    ときに「どのセクションにも出ない行」が静かに生まれるのを防ぐため。
    """
    name = str(label or "").split("　")[0].strip()
    return _PROJECT_SECTIONS.get(name, SECTION_UNKNOWN)


def case_label_of(label) -> str:
    """見出しラベルの「区分」部分(アドバリューの 8-1 等)。無ければ空文字。"""
    parts = str(label or "").split("　", 1)
    return parts[1].strip() if len(parts) > 1 else ""


# 会議用売上表では「その他」(単独チラシ)も案件ごとに行を分ける
# (依頼⑦・2026-08-27＝単独案件が複数あれば行を増やす)。
BULK_SPLIT_LABEL_PROJECTS = ("アドバリュー", "その他")


def build_bulk_rows(*, receivables, petty, payables, contract_lines, manual, id2proj, id2cat):
    """月次の全案件ぶんを、会議用売上表の行としてまとめて作る(依頼②・2026-08-20)。

    グルーピングは posting_logic.company_summary_by_project と同じ
    (アドバリューと「その他」を案件区分ごとに分ける)。戻り値は
    [(見出しラベル, build_row()の戻り値), ...]。ラベルは実物の
    「ｱﾄﾞ・バリュー　8-1」のような結合セル1つぶんの表示名。
    """
    from common import posting_logic as _pl

    summary = _pl.company_summary_by_project(
        receivables=receivables, payables=payables, petty=petty,
        contract_lines=contract_lines, manual=manual, id2proj=id2proj,
        split_labels_for=BULK_SPLIT_LABEL_PROJECTS)

    koutsuhi, nomimono = {}, {}
    for r in petty:
        pname = id2proj.get(r.get("project_id"), "")
        key = _pl.project_group_key(pname, r.get("other_label"), BULK_SPLIT_LABEL_PROJECTS)
        cat = id2cat.get(r.get("category_id"))
        amt = int(r.get("amount") or 0)
        if cat == "駐車場代":
            koutsuhi[key] = koutsuhi.get(key, 0) + amt
        elif cat == "飲み物代":
            nomimono[key] = nomimono.get(key, 0) + amt

    out = []
    for s in summary:
        key = (s["案件"], s["区分"])
        label = f'{s["案件"]}　{s["区分"]}' if s["区分"] else s["案件"]
        row = build_row(hakko_gou=None, ban_mei=None,
                        uriage_zeikomi=s["売上"], genka_goukei_zeikomi=s["原価"],
                        koutsuhi=koutsuhi.get(key, 0), nomimono=nomimono.get(key, 0))
        out.append((label, row))
    return out


def build_workbook(rows) -> bytes:
    """会議用売上表と同じ列見出しのExcelを bytes で返す(コピー＆ペースト用)。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "売上表用出力"
    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = openpyxl.styles.Font(bold=True)
    for i, r in enumerate(rows):
        for col, h in enumerate(HEADERS, start=1):
            ws.cell(row=2 + i, column=col, value=r.get(h))
    buf = io.BytesIO()
    wb.save(buf)
    return excel_io.freeze_xlsx_bytes(buf.getvalue())


def uriagehyo_filename() -> str:
    return "会議用売上表_出力.xlsx"
