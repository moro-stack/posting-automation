"""会議用売上表を、大阪支社からもらった新テンプレート
`templates/KPS(大阪)原価売上表テンプレート.xlsx` のレイアウトどおりに作る。

依頼⑦(2026-08-27 大橋様)。8/20 の実装は「実物と同じ書式の1行1案件の表」を
アプリが独自に組み立てるものだったが、実際の会議用売上表は

    タイトル
    見出し(2行)
    京阪南 / 京阪北            ← 版名(B列)を持つ2行
    (区切りの細い空行)
    単独チラシ                 ← 案件「その他」。件数ぶん行を増やす
    (区切りの細い空行)
    ｱﾄﾞ・バリュー 8-1〜8-5      ← 週ごとの5枠(月が変われば9-1〜9-5)
    (区切りの細い空行)
    リビング                   ← 2行(A列・B列は縦結合)
    (区切りの細い空行)
    集計(北大阪営業部/アド・バリュー/リビング・プロシード/計)

というセクション構成になっている。テンプレートの書式(フォント・罫線・列幅・
表示形式・行高)をそのまま複写し、アプリが持っている数字だけを流し込む。

🔴 行が増えても数式がズレないようにする(依頼⑦の例外対応)
単独チラシは案件ごとに行を増やす前提なので、テンプレートの数式
(=SUM(J4:J7) など)をそのままコピーすると参照が狂う。ここでは
**実際に書いた行番号から数式を組み立て直す**方式にした
(テンプレート側をテーブル化する案より、月ごとに行数が変わる実態に合う)。

🔴 J列(売上合計 税抜)は数式ではなく実額を書く
テンプレートの J は =D+F+H+I（ぱど/チラシ/仕分けの内訳の合計）だが、
アプリはその内訳(部数・区分別売上)を持っていない。数式のまま出すと
0円の表になってしまうため、アプリが持っている売上額をそのまま入れる。
実物の売上表もアドバリュー行は同じく実額を直接入れている。
C〜I(部数・内訳)と備考は空欄のまま出すので、会議前に手で足せる。
"""
import io
from copy import copy
from pathlib import Path

import openpyxl

from common import excel_io
from common import uriagehyo as U
from common.uriagehyo import HEADERS

TEMPLATE_PATH = (Path(__file__).resolve().parent.parent
                 / "templates" / "KPS(大阪)原価売上表テンプレート.xlsx")

TITLE_ROW = 1
HEADER_ROWS = (2, 3)
LAST_COL = 17          # A〜Q列(見出しの実在範囲)
TITLE_MERGE_END_COL = 15   # 実物は A1:O1
PRINT_LAST_COL = "O"

HEADER_MERGES = (
    "A2:A3", "B2:B3", "C2:D2", "E2:F2", "G2:H2",
    "I2:I3", "J2:J3", "K2:K3", "L2:L3", "M2:M3", "N2:N3", "Q2:Q3",
)

# テンプレートのどの行を、どの役割の見本(書式)にするか。
STYLE_ROWS = {
    "pado": 4,          # 京阪南/京阪北(A=発行号・B=版名)
    "spacer": 6,        # セクションの区切りの細い空行
    "tandoku": 7,       # 単独チラシ(A:B結合)
    "advalue": 9,       # ｱﾄﾞ・バリュー 週次(A:B結合)
    "living_top": 15,   # リビング 1行目
    "living_rest": 16,  # リビング 2行目以降
    "summary": 18,      # 集計の1行目
    "summary_mid": 19,
    "summary_last": 21,  # 「計」の行
    "note": 22,         # ※ 小数点以下、四捨五入
}

# テンプレートに元から入っている案内文(単独チラシの枠)。データが無いときはこれを残す。
TANDOKU_HINT = "単独チラシ(あれば入力、案件ごとに行増やす)"
ADVALUE_PREFIX = "ｱﾄﾞ・バリュー　"   # 実物の表記(半角カナ＋全角スペース)
ADVALUE_SLOTS = 5                    # テンプレートの週枠(8-1〜8-5)
LIVING_ROWS = 2                      # テンプレートのリビング枠
NOTE_TEXT = "※ 小数点以下、四捨五入"

# 集計行の備考(L列)。実物と同じ文言・先頭の半角スペースまで合わせる。
SUMMARY_LABELS = {
    "pado": " 関西ぱど",
    "unknown": " その他",
    "advalue": " アド・バリュー",
    "living": " リビング・プロシード",
}

# 列番号(1始まり)
COL_A, COL_B = 1, 2
COL_BUSUU_FIRST, COL_BUSUU_LAST = 3, 9    # C〜I(ぱど/チラシ/仕分け/その他)
COL_J, COL_K, COL_L = 10, 11, 12
COL_M, COL_N, COL_O, COL_P = 13, 14, 15, 16


class UriagehyoTemplate:
    """テンプレートを読んで、書式(フォント・罫線・列幅・行高・表示形式)を引けるようにする。"""

    def __init__(self, path=TEMPLATE_PATH):
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(f"売上表テンプレートがありません: {self.path}")
        self._wb = openpyxl.load_workbook(self.path)
        self._ws = self._wb[self._wb.sheetnames[0]]

    @property
    def ws(self):
        return self._ws

    def copy_column_widths(self, ws):
        for key, dim in self._ws.column_dimensions.items():
            if dim.width:
                ws.column_dimensions[key].width = dim.width

    def row_height(self, row):
        dim = self._ws.row_dimensions.get(row)
        return dim.height if dim else None

    def style_row(self, role):
        """役割に対応するテンプレートの行番号。"""
        return STYLE_ROWS[role] if isinstance(role, str) else int(role)


def _copy_cell_style(src, dst):
    """openpyxl のセル書式を複写する。Font/Border/Alignment は共有すると
    片方の変更が他方に及ぶため、必ず copy() してから代入する。"""
    dst.font = copy(src.font)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    if src.fill is not None and src.fill.fill_type == "solid":
        dst.fill = copy(src.fill)


def advalue_slot_labels(month, n=ADVALUE_SLOTS):
    """その月のアドバリュー週枠のラベル。8月なら 8-1〜8-5、10月なら 10-1〜10-5。"""
    from common import advalue

    return [advalue.week_label(month, w) for w in range(1, n + 1)]


def plan_sections(bulk_rows, month):
    """どのセクションに何行出すかを決める(純粋関数・テストしやすいように分けてある)。

    戻り値: {セクション名: [(A列に出すラベル, 行dict or None), ...]}
    ・京阪南/京阪北は必ず1行ずつ(データが無くても枠を出す)。
    ・単独チラシはデータの件数ぶん。0件ならテンプレートの案内文だけの1行。
    ・アドバリューはその月の 8-1〜8-5 の枠を必ず出し、枠に無いラベル
      (8-6・7-1 など)は後ろに足す＝捨てない。
    ・リビングは最低2行(テンプレートの枠)。
    ・どのセクションにも当てはまらない案件は「未分類」に集める。
    """
    from common import advalue

    buckets = {U.SECTION_PADO_MINAMI: [], U.SECTION_PADO_KITA: [],
               U.SECTION_TANDOKU: [], U.SECTION_ADVALUE: [],
               U.SECTION_LIVING: [], U.SECTION_UNKNOWN: []}
    for label, row in bulk_rows or []:
        buckets[U.section_of(label)].append((label, row))

    plan = {}
    plan[U.SECTION_PADO_MINAMI] = [("京阪南", buckets[U.SECTION_PADO_MINAMI][0][1]
                                    if buckets[U.SECTION_PADO_MINAMI] else None)]
    plan[U.SECTION_PADO_KITA] = [("京阪北", buckets[U.SECTION_PADO_KITA][0][1]
                                  if buckets[U.SECTION_PADO_KITA] else None)]

    plan[U.SECTION_TANDOKU] = ([(label, row) for label, row in buckets[U.SECTION_TANDOKU]]
                               or [(TANDOKU_HINT, None)])

    by_case = {U.case_label_of(label): row for label, row in buckets[U.SECTION_ADVALUE]}
    slots = advalue_slot_labels(month)
    adv = [(ADVALUE_PREFIX + s, by_case.pop(s, None)) for s in slots]
    # 枠に無い週(8-6)・別の月のラベル(7-1)も行として残す
    for case in advalue.sort_week_labels(list(by_case.keys())):
        adv.append((ADVALUE_PREFIX + case, by_case[case]))
    plan[U.SECTION_ADVALUE] = adv

    living = [("リビング", row) for _, row in buckets[U.SECTION_LIVING]]
    while len(living) < LIVING_ROWS:
        living.append(("リビング", None))
    plan[U.SECTION_LIVING] = living

    plan[U.SECTION_UNKNOWN] = list(buckets[U.SECTION_UNKNOWN])
    return plan


def _write_data_row(ws, tpl, r, *, role, a_value, b_value, merge_ab, row):
    """データ行を1行書く。値の無い列はテンプレートの書式だけ入れて空欄にする。"""
    src_row = tpl.style_row(role)
    for c in range(1, LAST_COL + 1):
        _copy_cell_style(tpl.ws.cell(src_row, c), ws.cell(r, c))

    ws.cell(r, COL_A, a_value)
    if not merge_ab:
        ws.cell(r, COL_B, b_value)

    if row is not None:
        ws.cell(r, COL_J, row.get("売上合計（税抜）"))
        ws.cell(r, COL_M, row.get("交通費（駐車場代含）"))
        ws.cell(r, COL_N, row.get("飲み物"))
        ws.cell(r, COL_O, row.get("配布原価（税込）"))
    # 税込・原価合計はテンプレートと同じ数式で持たせる(手で数字を足しても追随する)
    ws.cell(r, COL_K, f"=ROUND(J{r}*1.1,0)")
    ws.cell(r, COL_P, f"=M{r}+N{r}+O{r}")

    if merge_ab:
        ws.merge_cells(start_row=r, start_column=COL_A, end_row=r, end_column=COL_B)
    h = tpl.row_height(src_row)
    if h:
        ws.row_dimensions[r].height = h


def _write_spacer(ws, tpl, r):
    src_row = tpl.style_row("spacer")
    for c in range(1, LAST_COL + 1):
        _copy_cell_style(tpl.ws.cell(src_row, c), ws.cell(r, c))
    h = tpl.row_height(src_row)
    if h:
        ws.row_dimensions[r].height = h


def _sum_formula(letter, rows):
    """指定した行だけを足す数式。行が飛んでいても正しく足せるように
    連番なら SUM(範囲)、飛んでいれば足し算にする。"""
    if not rows:
        return None
    if rows == list(range(rows[0], rows[-1] + 1)):
        return f"=SUM({letter}{rows[0]}:{letter}{rows[-1]})"
    return "=" + "+".join(f"{letter}{r}" for r in rows)


def _span_formula(letter, rows):
    """先頭行から最終行までを1つの範囲で足す数式。あいだの区切り行は空欄なので
    合計は変わらない(実物の総計 =SUM(M4:M20) と同じ書き方)。"""
    if not rows:
        return None
    return f"=SUM({letter}{rows[0]}:{letter}{rows[-1]})"


def _write_summary_row(ws, tpl, r, *, role, b_value, note, data_rows, cost_rows=None):
    src_row = tpl.style_row(role)
    for c in range(1, LAST_COL + 1):
        _copy_cell_style(tpl.ws.cell(src_row, c), ws.cell(r, c))
    if b_value is not None:
        ws.cell(r, COL_B, b_value)
    if note is not None:
        ws.cell(r, COL_L, note)
    for c in range(COL_BUSUU_FIRST, COL_BUSUU_LAST + 1):
        ws.cell(r, c, _sum_formula(openpyxl.utils.get_column_letter(c), data_rows))
    ws.cell(r, COL_J, _sum_formula("J", data_rows))
    ws.cell(r, COL_K, f"=ROUND(J{r}*1.1,0)")
    # 交通費・飲み物・配布原価も足し上げる。実物は月によって入れたり入れなかったり
    # だが(2026年3月度は入れている)、入れておかないと原価合計が常に0になる。
    for letter, col in (("M", COL_M), ("N", COL_N), ("O", COL_O)):
        if cost_rows is not None:
            # 総計行。実物と同じく先頭〜最終データ行を1つの範囲で足す
            ws.cell(r, col, _span_formula(letter, cost_rows))
        else:
            ws.cell(r, col, _sum_formula(letter, data_rows))
    ws.cell(r, COL_P, f"=M{r}+N{r}+O{r}")
    h = tpl.row_height(src_row)
    if h:
        ws.row_dimensions[r].height = h


def build_monthly_workbook(bulk_rows, year, month, *, tpl=None) -> bytes:
    """bulk_rows: uriagehyo.build_bulk_rows() の戻り値
    [(見出しラベル, build_row()の戻り値), ...] を、新テンプレートの
    レイアウトどおりに1シートへ書き出す。"""
    tpl = tpl or UriagehyoTemplate()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{year}年{month}月度 売上"

    tpl.copy_column_widths(ws)

    # ---- タイトル行 ----
    t = ws.cell(TITLE_ROW, 1, f"㈱ケイピーエス　大阪支社　　{year}年 {month}月度 売上表")
    _copy_cell_style(tpl.ws.cell(TITLE_ROW, 1), t)
    ws.merge_cells(start_row=TITLE_ROW, start_column=1,
                   end_row=TITLE_ROW, end_column=TITLE_MERGE_END_COL)
    th = tpl.row_height(TITLE_ROW)
    if th:
        ws.row_dimensions[TITLE_ROW].height = th

    # ---- 見出し2行(結合セル込み) ----
    for r in HEADER_ROWS:
        for c in range(1, LAST_COL + 1):
            src = tpl.ws.cell(r, c)
            _copy_cell_style(src, ws.cell(r, c, src.value))
        hh = tpl.row_height(r)
        if hh:
            ws.row_dimensions[r].height = hh
    for rng in HEADER_MERGES:
        ws.merge_cells(rng)

    plan = plan_sections(bulk_rows, month)
    r = HEADER_ROWS[-1] + 1
    used = {}          # セクション名 -> 書いた行番号のリスト

    def _block(section, *, role, merge_ab, b_from_label=False):
        nonlocal r
        rows = []
        for label, row in plan[section]:
            _write_data_row(ws, tpl, r, role=role,
                            a_value=None if b_from_label else label,
                            b_value=label if b_from_label else None,
                            merge_ab=merge_ab, row=row)
            rows.append(r)
            r += 1
        used[section] = rows
        return rows

    # 京阪南・京阪北(A=発行号は空欄・B=版名)
    _block(U.SECTION_PADO_MINAMI, role="pado", merge_ab=False, b_from_label=True)
    _block(U.SECTION_PADO_KITA, role="pado", merge_ab=False, b_from_label=True)
    _write_spacer(ws, tpl, r); r += 1

    _block(U.SECTION_TANDOKU, role="tandoku", merge_ab=True)
    _write_spacer(ws, tpl, r); r += 1

    if plan[U.SECTION_UNKNOWN]:
        _block(U.SECTION_UNKNOWN, role="tandoku", merge_ab=True)
        _write_spacer(ws, tpl, r); r += 1
    else:
        used[U.SECTION_UNKNOWN] = []

    _block(U.SECTION_ADVALUE, role="advalue", merge_ab=True)
    _write_spacer(ws, tpl, r); r += 1

    living_rows = []
    for i, (label, row) in enumerate(plan[U.SECTION_LIVING]):
        _write_data_row(ws, tpl, r, role="living_top" if i == 0 else "living_rest",
                        a_value=None, b_value=label if i == 0 else None,
                        merge_ab=False, row=row)
        living_rows.append(r)
        r += 1
    used[U.SECTION_LIVING] = living_rows
    if len(living_rows) > 1:
        # 発行号(A)・版名(B)はリビングの行数ぶん縦結合(実物と同じ)
        ws.merge_cells(start_row=living_rows[0], start_column=COL_A,
                       end_row=living_rows[-1], end_column=COL_A)
        ws.merge_cells(start_row=living_rows[0], start_column=COL_B,
                       end_row=living_rows[-1], end_column=COL_B)
    _write_spacer(ws, tpl, r); r += 1

    # ---- 集計ブロック ----
    # 単独チラシは実物どおり「北大阪営業部(関西ぱど)」に含める。
    pado_rows = (used[U.SECTION_PADO_MINAMI] + used[U.SECTION_PADO_KITA]
                 + used[U.SECTION_TANDOKU])
    summary_top = r
    _write_summary_row(ws, tpl, r, role="summary", b_value="北大阪営業部",
                       note=SUMMARY_LABELS["pado"], data_rows=pado_rows)
    r += 1
    if used[U.SECTION_UNKNOWN]:
        _write_summary_row(ws, tpl, r, role="summary_mid", b_value="その他",
                           note=SUMMARY_LABELS["unknown"],
                           data_rows=used[U.SECTION_UNKNOWN])
        r += 1
    adv_summary_row = r
    _write_summary_row(ws, tpl, r, role="summary_mid", b_value=None,
                       note=SUMMARY_LABELS["advalue"], data_rows=used[U.SECTION_ADVALUE])
    r += 1
    _write_summary_row(ws, tpl, r, role="summary_mid", b_value=None,
                       note=SUMMARY_LABELS["living"], data_rows=used[U.SECTION_LIVING])
    living_summary_row = r
    r += 1

    all_data_rows = sorted(sum((used[s] for s in used), []))
    subtotal_rows = list(range(summary_top, r))
    _write_summary_row(ws, tpl, r, role="summary_last", b_value="計", note=None,
                       data_rows=subtotal_rows, cost_rows=all_data_rows)
    total_row = r
    r += 1

    # 実物は「アド・バリュー」「リビング・プロシード」の版名欄(B)をまとめて結合している
    if living_summary_row > adv_summary_row:
        ws.merge_cells(start_row=adv_summary_row, start_column=COL_B,
                       end_row=living_summary_row, end_column=COL_B)
    # 月度(A列)は集計ブロック全体で縦結合
    a = ws.cell(summary_top, COL_A, f"{month}月度")
    _copy_cell_style(tpl.ws.cell(tpl.style_row("summary"), COL_A), a)
    ws.merge_cells(start_row=summary_top, start_column=COL_A,
                   end_row=total_row, end_column=COL_A)

    # ---- 注記 ----
    note_src = tpl.style_row("note")
    for c in range(1, LAST_COL + 1):
        _copy_cell_style(tpl.ws.cell(note_src, c), ws.cell(r, c))
    ws.cell(r, COL_B, NOTE_TEXT)
    nh = tpl.row_height(note_src)
    if nh:
        ws.row_dimensions[r].height = nh

    ws.print_area = f"A1:{PRINT_LAST_COL}{total_row}"
    ws.page_setup.orientation = tpl.ws.page_setup.orientation or "landscape"
    ws.page_setup.paperSize = tpl.ws.page_setup.paperSize
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(
        fitToPage=True)

    buf = io.BytesIO()
    wb.save(buf)
    return excel_io.freeze_xlsx_bytes(buf.getvalue())


def monthly_filename(year, month) -> str:
    return f"{year}年{month}月度_売上表.xlsx"
