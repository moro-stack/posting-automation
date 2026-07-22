"""現行 .xlsm の あてがみ シートを、データを消したテンプレとして切り出す。

実データ(実在の配送物名・担当者名・地区コード)を残さないため可変セルをクリアする。
固定文字 A4「担当地区　：　」/ C4=0 / J2「 部」はテンプレの値を残す。
"""
import os
import openpyxl

XLSM = r"C:\Users\moro\Downloads\@マクロ（京阪北版）あて紙　マスタ (1).xlsm"
OUT = os.path.join(os.path.dirname(__file__), "..", "templates", "あて紙テンプレート_京阪.xlsx")

wb = openpyxl.load_workbook(XLSM)  # 書式・結合・印刷設定を保持（read_only不可）
ws = wb["あてがみ"]

# 可変セルをクリア（固定文字 A4/C4/J2 は残す）
for coord in ["A1", "A3", "D4", "G2", "H18", "I18"]:
    ws[coord] = None
for r in range(5, 19):          # B5:J18 の明細領域
    for c in range(2, 11):
        ws.cell(row=r, column=c).value = None

# あてがみ 以外のシートを削除
for name in list(wb.sheetnames):
    if name != "あてがみ":
        del wb[name]
ws.title = "あて紙テンプレート"

os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)
wb.save(os.path.abspath(OUT))
print("WROTE", os.path.abspath(OUT))
