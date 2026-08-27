import io

import openpyxl

from common import advalue as V


def test_parse_choume_edge_cases():
    # 現行マクロが取りこぼしていたケースを含む
    assert V.parse_choume_from_gaiku("1丁目1～21") == {1}
    assert V.parse_choume_from_gaiku("2・3丁目") == {2, 3}
    assert V.parse_choume_from_gaiku("4・5丁目") == {4, 5}
    assert V.parse_choume_from_gaiku("6丁目51～57・7丁目") == {6, 7}     # 番地57を丁目に誤認しない
    assert V.parse_choume_from_gaiku("1丁目・2丁目11～34") == {1, 2}      # 2丁目を落とさない
    assert V.parse_choume_from_gaiku("１丁目") == {1}                     # 全角
    assert V.parse_choume_from_gaiku("") == set()                        # 丁目なし=ワイルドカード


def test_split_town_choume():
    assert V.split_town_choume("大久保町1") == ("大久保町", 1)
    assert V.split_town_choume("大久保町1丁目") == ("大久保町", 1)
    assert V.split_town_choume("宮前町") == ("宮前町", None)


def test_match_area_with_choume_range():
    # 大久保町1丁目が 913101(1～21) と 913102(22～52) に分かれている → 依頼表「大久保町1」は両方に被り
    index = V.build_minami_index([
        ["配送管理表"],
        ["配布日", "号数", "ルート", "異動", "配送順位", "ぱどんな", "住所", "電話番号",
         "担当地区", "チラシコード", "配送物", "配布部数", "配送備考", "町界名",
         "街区（番地）名称", "受注種別", "チラシサイズ"],
        [None, None, None, None, None, None, None, None, "913101", None, "91 ぱど",
         438, None, "大久保町", "1丁目1～21"],
        [None, None, None, None, None, None, None, None, "913102", None, "91 ぱど",
         438, None, "大久保町", "1丁目22～52"],
        [None, None, None, None, None, None, None, None, "923301", None, "91 ぱど",
         110, None, "大庭町", "1・2丁目"],
    ])
    assert V.match_area("大久保町", 1, index) == ["913101", "913102"]
    assert V.match_area("大庭町", 1, index) == ["923301"]
    assert V.match_area("大庭町", 2, index) == ["923301"]
    assert V.match_area("存在しない町", 1, index) == []     # 未マッチ


def test_build_report_rows_and_summary():
    index = V.build_minami_index([
        ["配送管理表"],
        ["配布日", "号数", "ルート", "異動", "配送順位", "ぱどんな", "住所", "電話番号",
         "担当地区", "チラシコード", "配送物", "配布部数", "配送備考", "町界名",
         "街区（番地）名称", "受注種別", "チラシサイズ"],
        [None] * 8 + ["913101", None, "x", 1, None, "大久保町", "1丁目"],
    ])
    irai = [
        {"geocode": "27209001001", "shiku": "守口市", "choume_name": "大久保町1",
         "setai": 840, "kyodo_setai": 134, "zantei": 570, "client1": "×", "client2": "○",
         "haifubi_memo": "6月9日", "heihai": 570, "biko": "サンプル店", "toukin": "なし",
         "haifubi": "6/8-6/14"},
        {"geocode": "27209099999", "shiku": "守口市", "choume_name": "被らない町9",
         "setai": 100, "kyodo_setai": 10, "zantei": 100, "client1": "×", "client2": "○",
         "haifubi_memo": "", "heihai": 100, "biko": "店", "toukin": "なし",
         "haifubi": "6/8-6/14"},
    ]
    rows = V.build_report_rows(irai, index)
    assert rows[0]["担当地区(被り)"] == "913101"
    assert rows[0]["被り"] == "被り"
    assert rows[0]["世帯数15"] == 840
    assert rows[1]["担当地区(被り)"] == "未マッチ"
    assert rows[1]["被り"] == "非被り"
    assert V.overlap_summary(rows) == {"total": 2, "overlap": 1, "non_overlap": 1}


def test_build_report_workbook():
    rows = [{"町丁目名": "サンプル町1", "市区名": "サンプル市", "担当地区(被り)": "913101",
             "被り": "被り", "世帯数15": 840, "暫定部数": 570, "クライアント①": "×",
             "クライアント②": "○", "配布日メモ": "6月9日", "併配部数": 570, "店舗名": "サンプル店",
             "投禁物件": "なし", "配布日": "6/8-6/14"}]
    data = V.build_advalue_report_workbook(rows)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "町丁目名"
    assert ws["C1"].value == "担当地区(被り)"
    assert ws["G1"].value == "クライアント①"
    assert ws["A2"].value == "サンプル町1"
    assert ws["C2"].value == "913101"
    assert ws["D2"].value == "被り"


# ===== 週ごとの案件(8-1/8-2/8-3等)を分けて管理したい（依頼⑥・2026-08-19大橋様ご指摘） =====


def test_build_advalue_report_workbook_without_case_label_is_unchanged():
    """case_label を渡さなければ、これまでどおりヘッダーがA1から始まる。"""
    rows = [{"町丁目名": "サンプル町1", "市区名": "サンプル市", "担当地区(被り)": "913101",
             "被り": "被り", "世帯数15": 840, "暫定部数": 570, "クライアント①": "×",
             "クライアント②": "○", "配布日メモ": "6月9日", "併配部数": 570, "店舗名": "サンプル店",
             "投禁物件": "なし", "配布日": "6/8-6/14"}]
    data = V.build_advalue_report_workbook(rows)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "町丁目名"


def test_build_advalue_report_workbook_with_case_label_adds_title_row():
    """アドバリューは週ごとに『8-1』『8-2』『8-3』のように案件を分けて管理しているため、
    ラベルを渡したときは報告書に見出し行として入れ、ヘッダーは1行下にずれる。"""
    rows = [{"町丁目名": "サンプル町1", "市区名": "サンプル市", "担当地区(被り)": "913101",
             "被り": "被り", "世帯数15": 840, "暫定部数": 570, "クライアント①": "×",
             "クライアント②": "○", "配布日メモ": "6月9日", "併配部数": 570, "店舗名": "サンプル店",
             "投禁物件": "なし", "配布日": "6/8-6/14"}]
    data = V.build_advalue_report_workbook(rows, case_label="8-1")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "アドバリュー エリア被り報告書　8-1"
    assert ws["A2"].value == "町丁目名"
    assert ws["A3"].value == "サンプル町1"


def test_advalue_filename_without_case_label():
    assert V.advalue_filename() == "アドバリュー_エリア被り報告書.xlsx"


def test_advalue_filename_with_case_label():
    """ファイル名にも案件を入れる。週ごとに生成しても上書きせず分けて保存できるように。"""
    assert V.advalue_filename("8-1") == "アドバリュー_エリア被り報告書_8-1.xlsx"


def test_advalue_filename_strips_characters_unsafe_for_filenames():
    """案件名に / などが入っても壊れたファイル名にならないこと。"""
    assert "/" not in V.advalue_filename("8/1")


# ===== 週次の案件区分（依頼⑥・2026-08-27 大橋様ご指摘） =====
# アドバリューは週ごとに案件が来るので、登録時に必ず週を選ばせる。
# 月をまたいだら自動で「9-1」「10-1」になること。


def test_week_choices_are_month_independent():
    """画面で選ぶのは月に依存しない『1週目〜5週目』。
    月の部分は登録する伝票自身の日付から決めるので、9月になれば黙って
    9-1 になり、10月なら 10-1 になる(選択肢を作り直す必要が無い)。"""
    assert V.WEEK_CHOICES == ("1週目", "2週目", "3週目", "4週目", "5週目")


def test_week_label_builds_the_month_dash_week_form():
    assert V.week_label(8, 1) == "8-1"
    assert V.week_label(9, 3) == "9-3"
    assert V.week_label(10, 5) == "10-5"


def test_week_index_from_choice():
    assert V.week_index("1週目") == 1
    assert V.week_index("5週目") == 5
    assert V.week_index("") is None
    assert V.week_index(None) is None
    assert V.week_index("（選択してください）") is None


def test_month_of_reads_a_date_string():
    assert V.month_of("2026-09-03") == 9
    assert V.month_of("2026-10") == 10


def test_month_of_reads_a_date_object():
    import datetime
    assert V.month_of(datetime.date(2026, 10, 1)) == 10


def test_month_of_falls_back_to_today_when_the_date_is_missing_or_broken():
    """日付が未入力・読めないときに登録を止めるのはやり過ぎなので、
    登録日の月に倒す(あとから号別明細で直せる)。"""
    assert V.month_of("", today="2026-09-15") == 9
    assert V.month_of(None, today="2026-09-15") == 9
    assert V.month_of("ぐちゃぐちゃ", today="2026-11-01") == 11


def test_week_label_for_uses_the_month_of_the_voucher_not_of_today():
    """🔴 9月の伝票を10月に入力しても「9-1」。登録日ではなく伝票の日付で決める。"""
    assert V.week_label_for("2026-09-04", "1週目", today="2026-10-20") == "9-1"


def test_week_label_for_rolls_over_the_month_automatically():
    assert V.week_label_for("2026-08-04", "1週目") == "8-1"
    assert V.week_label_for("2026-09-01", "1週目") == "9-1"
    assert V.week_label_for("2026-10-01", "1週目") == "10-1"


def test_week_label_for_returns_none_when_no_week_is_chosen():
    """週が未選択なら区分は作らない(呼び出し側が『選んでください』を出す)。"""
    assert V.week_label_for("2026-08-04", None) is None
    assert V.week_label_for("2026-08-04", "（選択してください）") is None


def test_parse_week_label_reads_existing_labels():
    """8/20以前に手入力された「8-1」もそのまま読めること(過去データ互換)。"""
    assert V.parse_week_label("8-1") == (8, 1)
    assert V.parse_week_label("10-5") == (10, 5)
    assert V.parse_week_label("配夢") is None
    assert V.parse_week_label(None) is None


def test_sort_week_labels_orders_by_month_then_week_not_by_string():
    """🔴 文字列順だと「10-1」が「8-1」より前に来て、月をまたぐと並びが壊れる。"""
    assert V.sort_week_labels(["10-1", "8-2", "9-1", "8-1"]) == [
        "8-1", "8-2", "9-1", "10-1"]


def test_sort_week_labels_keeps_non_week_labels_at_the_end():
    """アドバリュー以外の使い方(手入力の案件名)が混ざっても落とさない。"""
    assert V.sort_week_labels(["9-1", "配夢", "8-1"]) == ["8-1", "9-1", "配夢"]


def test_week_display_is_readable_in_the_issue_page():
    assert V.week_display("8-1") == "8月 1週目"
    assert V.week_display("10-5") == "10月 5週目"
    assert V.week_display("配夢") == "配夢"     # 週でないものはそのまま


# ----- 案件区分の決定（登録画面が使う唯一の入口） -----


def test_resolve_case_label_requires_a_week_for_advalue():
    """🔴 依頼⑥の肝: アドバリューは週を選ばないと登録させない。"""
    label, err = V.resolve_case_label(
        "アドバリュー", date_value="2026-08-04", week_choice=None, free_text="")
    assert label is None
    assert err and "週" in err


def test_resolve_case_label_builds_the_week_label_for_advalue():
    label, err = V.resolve_case_label(
        "アドバリュー", date_value="2026-09-04", week_choice="2週目", free_text="")
    assert (label, err) == ("9-2", None)


def test_resolve_case_label_ignores_free_text_for_advalue():
    """アドバリューでは自由入力ではなく必ず週から作る(表記ゆれを作らない)。"""
    label, _ = V.resolve_case_label(
        "アドバリュー", date_value="2026-08-04", week_choice="1週目", free_text="八月一週")
    assert label == "8-1"


def test_resolve_case_label_uses_free_text_for_sonota():
    label, err = V.resolve_case_label(
        "その他", date_value="2026-08-04", week_choice=None, free_text=" 買取専科 ")
    assert (label, err) == ("買取専科", None)


def test_resolve_case_label_sonota_without_text_is_none_but_not_an_error():
    """「その他」の案件名は今までどおり任意(空でも登録できる)。"""
    assert V.resolve_case_label(
        "その他", date_value="2026-08-04", week_choice=None, free_text="") == (None, None)


def test_resolve_case_label_is_none_for_other_projects():
    for name in ("関西ぱど：京阪北版", "リビングプロシード", "(なし)", None):
        assert V.resolve_case_label(
            name, date_value="2026-08-04", week_choice="1週目", free_text="x") == (None, None)
