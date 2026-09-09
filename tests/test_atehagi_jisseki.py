import io

import openpyxl

from common import atehagi as A
from tests.test_atehagi import _sample_table


def test_jisseki_filename():
    rows = A.rows_from_table(_sample_table())
    name = A.jisseki_filename(A.KEIHAN_KITA, rows)
    assert name.startswith("京阪北版_挟み込み実績表_1248号_")
    assert name.endswith(".xlsx")


def test_jisseki_courses_basic():
    groups = A.group_by_chiku(A.rows_from_table(_sample_table()))
    courses = A.jisseki_courses(groups, A.KEIHAN_KITA)
    # 10101 は チラシ有り→載る。46501 は ぱどのみ→除外。
    assert len(courses) == 1
    c = courses[0]
    assert c["code"] == "010101"
    assert c["chiku_name"] == "枚方・交野"
    assert c["course_name"] == "010101 枚方・交野"
    assert c["flyers"] == [{"name": "サンプル生協", "count": 224}]


def _multi_flyer_table():
    header = ["配布日", "号数", "ルート", "異動", "配送順位", "ぱどんな", "住所",
              "電話番号", "担当地区", "チラシコード", "配送物", "配布部数", "配送備考",
              "町界名", "街区（番地）名称", "受注種別", "チラシサイズ"]
    return [
        ["配送管理表", "2026-06-26", "(1248号)", "作成日:2026/06/18"],
        header,
        ["2026-06-26", 1248, 0, None, None, "テスト太郎", None, None, 10101, None,
         "01 ぱど", 400, " ", "大住ヶ丘", "1丁目", None, None],
        ["2026-06-26", 1248, 0, None, None, "テスト太郎", None, None, 10101, 400,
         "SUUMO/注文住宅", 400, " ", "大住ヶ丘", "1丁目", "全戸配布(チラシ)", "Ｂ４"],
        ["2026-06-26", 1248, 0, None, None, "テスト太郎", None, None, 10101, 400,
         "ECC/宮野", 400, " ", "大住ヶ丘", "1丁目", "全戸配布(チラシ)", "Ｂ４"],
    ]


def test_jisseki_courses_multi_flyer_keeps_order():
    groups = A.group_by_chiku(A.rows_from_table(_multi_flyer_table()))
    courses = A.jisseki_courses(groups, A.KEIHAN_KITA)
    assert len(courses) == 1
    assert courses[0]["flyers"] == [
        {"name": "SUUMO/注文住宅", "count": 400},
        {"name": "ECC/宮野", "count": 400},
    ]


# ===== 通し番号（依頼①・2026-09-04大橋様ご依頼） =====


def test_circled_number_within_range():
    assert A.circled_number(1) == "①"
    assert A.circled_number(9) == "⑨"
    assert A.circled_number(20) == "⑳"


def test_circled_number_falls_back_beyond_20():
    assert A.circled_number(21) == "(21)"


def test_assign_course_numbers_restarts_at_1_per_leader():
    courses = [
        {"code": "1", "leader": "山田"},
        {"code": "2", "leader": "山田"},
        {"code": "3", "leader": "外注先"},
        {"code": "4", "leader": "山田"},
    ]
    numbers = A.assign_course_numbers(courses)
    assert numbers == {"1": 1, "2": 2, "3": 1, "4": 3}


def test_assign_course_numbers_is_pure_and_repeatable():
    """通し番号はコースの並び順だけで決まる(同じ入力なら何度呼んでも同じ)。
    あて紙・実績表それぞれ別のボタンから独立に呼んでも一致させるための前提。"""
    courses = [{"code": "1", "leader": "山田"}, {"code": "2", "leader": "山田"}]
    assert A.assign_course_numbers(courses) == A.assign_course_numbers(list(courses))


def test_md_from():
    assert A._md_from("2026-06-26") == "6/26"
    import datetime
    assert A._md_from(datetime.date(2026, 6, 26)) == "6/26"
    assert A._md_from("") == ""


def test_build_daishi_title_and_first_course():
    groups = A.group_by_chiku(A.rows_from_table(_sample_table()))
    courses = A.jisseki_courses(groups, A.KEIHAN_KITA)
    data = A.build_jisseki_daishi_workbook(courses, A.KEIHAN_KITA, gou=1248,
                                           haifubi="2026-06-26")
    import io
    wb = openpyxl.load_workbook(io.BytesIO(data))
    # 🔴 2026-08-27 大橋様ご指摘: 外注先/リーダーごとに別のシート(タブ)に分ける。
    # タブ名がその外注先/リーダー名そのものになる。
    ws = wb["テスト太郎"]
    # 🔴 2026-09-04大橋様追加ご依頼: エリア名はコース見出しから見出し行(A1・リーダー名の
    # 横)へ移した。コース見出し(A2)は通し番号＋担当地区コード(6桁)のみにする。
    assert ws["A1"].value == "6/26 ／ 1248号　京阪北版　テスト太郎　枚方・交野"
    assert ws["A2"].value == "1\n010101"                # ヘッダー(通し番号＋担当地区コード)
    assert ws["A3"].value == "①224"                     # 本文(①チラシ種類数＋ぱど本誌部数)
    assert ws["A4"].value == "サイン："                # サイン行
    assert ws["A4"].border.bottom.style is not None    # 署名用の横線
    assert ws.row_dimensions[4].height == 40            # サインの手書きスペースを広く
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.paperSize == 8                  # A3


# ===== リーダー名・外注先ごとの分割表示（依頼④・2026-08-19 大橋様ご指摘） =====


def test_jisseki_courses_includes_leader():
    groups = A.group_by_chiku(A.rows_from_table(_sample_table()))
    courses = A.jisseki_courses(groups, A.KEIHAN_KITA)
    assert courses[0]["leader"] == "テスト太郎"


def _leader_courses():
    return [
        {"code": "010101", "chiku_name": "A地区", "course_name": "010101 A地区",
         "flyers": [{"name": "f1", "count": 100}], "leader": "山田"},
        {"code": "010102", "chiku_name": "B地区", "course_name": "010102 B地区",
         "flyers": [{"name": "f2", "count": 200}], "leader": "外注先会社"},
        {"code": "010103", "chiku_name": "C地区", "course_name": "010103 C地区",
         "flyers": [{"name": "f3", "count": 300}], "leader": "山田"},
    ]


def _wb(courses, **kw):
    data = A.build_jisseki_daishi_workbook(courses, A.KEIHAN_KITA, gou=1,
                                           haifubi="2026-06-26", **kw)
    return openpyxl.load_workbook(io.BytesIO(data))


def test_build_daishi_splits_leaders_into_separate_sheets():
    """🔴 依頼④(2026-08-27 大橋様): 印刷して使うので、外注先/リーダーごとに
    別々のシート(タブ)に分ける。タブ名はその名前そのもの。"""
    wb = _wb(_leader_courses())
    assert wb.sheetnames == ["山田", "外注先会社"]   # 初出順


def test_build_daishi_each_sheet_holds_only_that_leaders_courses():
    wb = _wb(_leader_courses())
    yamada = wb["山田"]
    # 離れて出てきた同じリーダーのコースも同じシートにまとまる。
    # 🔴 2026-09-04: 1コース=1列になったので、2件目は隣の列(B)。
    # 通し番号はリーダーが変わるごとに1から(山田の2件目は2)。
    # コース見出しは通し番号＋担当地区コードのみ(エリア名は見出し行へ移した)。
    assert yamada["A2"].value == "1\n010101"
    assert yamada["B2"].value == "2\n010103"
    gaichu = wb["外注先会社"]
    assert gaichu["A2"].value == "1\n010102"
    # 他人のコースが混ざっていないこと(担当地区コードで判定)
    values = {c.value for row in gaichu.iter_rows() for c in row}
    assert not any(v and ("010101" in str(v) or "010103" in str(v)) for v in values)


def test_build_daishi_loses_no_course_when_splitting():
    """🔴 取りこぼしゼロ。分割した結果どのタブにも出ないコースが出ないこと。"""
    courses = _leader_courses()
    numbers = A.assign_course_numbers(courses)
    expected = {f'{numbers[c["code"]]}\n{c["code"]}' for c in courses}
    wb = _wb(courses)
    seen = set()
    for name in wb.sheetnames:
        for row in wb[name].iter_rows():
            for c in row:
                if c.value in expected:
                    assert c.value not in seen, f"{c.value} が2か所に出ている"
                    seen.add(c.value)
    assert seen == expected


def test_build_daishi_puts_the_leader_name_and_area_in_the_title_of_each_sheet():
    """印刷したときに誰の紙か分かるよう、見出しにも名前を出す。
    🔴 2026-09-04大橋様追加ご依頼: エリア名はここ(リーダー名の横)に出す。"""
    wb = _wb(_leader_courses())
    assert wb["山田"]["A1"].value == "6/26 ／ 1号　京阪北版　山田　A地区／C地区"
    assert wb["外注先会社"]["A1"].value == "6/26 ／ 1号　京阪北版　外注先会社　B地区"


def test_build_daishi_sets_print_settings_on_every_sheet():
    """タブごとに印刷するので、印刷設定は全シートに要る。"""
    wb = _wb(_leader_courses())
    for name in wb.sheetnames:
        ws = wb[name]
        assert ws.page_setup.orientation == "landscape"
        assert ws.print_area, f"{name} に印刷範囲が無い"


def test_build_daishi_sheet_titles_use_the_shared_sanitizer():
    """🔴 シート名の禁止文字・31文字制限は excel_io の共通処理に任せる
    (2026-08-10 の仕分け表と同じ事故＝Excelの修復を繰り返さないため)。"""
    from common import excel_io as X

    ng = set(X.EXCEL_NG_SHEET_CHARS)
    courses = [
        {"code": "1", "chiku_name": "A", "course_name": "c1",
         "flyers": [{"name": "f", "count": 1}], "leader": "メイト／南川"},
        {"code": "2", "chiku_name": "B", "course_name": "c2",
         "flyers": [{"name": "f", "count": 1}], "leader": "あ" * 40},
    ]
    wb = _wb(courses)
    for name in wb.sheetnames:
        assert not set(name) & ng, f"{name!r} に禁止文字が残る"
        assert len(name) <= 31, f"{name!r} が31文字を超える"


def test_build_daishi_dedupes_sheet_titles_that_collide_after_trimming():
    """31文字で切ると同じ名前になる外注先が2つあっても、シートが消えないこと。"""
    courses = [
        {"code": "1", "chiku_name": "A", "course_name": "c1",
         "flyers": [{"name": "f", "count": 1}], "leader": "あ" * 31 + "1"},
        {"code": "2", "chiku_name": "B", "course_name": "c2",
         "flyers": [{"name": "f", "count": 1}], "leader": "あ" * 31 + "2"},
    ]
    wb = _wb(courses)
    assert len(wb.sheetnames) == 2
    assert len(set(wb.sheetnames)) == 2


def test_build_daishi_without_leader_key_keeps_one_sheet():
    """leader を持たないコース(既存の呼び出し方)は、以前どおり1枚の実績表。"""
    courses = [
        {"code": "010101", "chiku_name": "A地区", "course_name": "010101 A地区",
         "flyers": [{"name": "f1", "count": 100}]},
    ]
    wb = _wb(courses)
    assert wb.sheetnames == ["実績表"]
    ws = wb["実績表"]
    # 名前(リーダー)が無いのでタブ見出しには付けないが、エリア名は付く。
    assert ws["A1"].value == "6/26 ／ 1号　京阪北版　A地区"
    assert ws["A2"].value == "1\n010101"


def test_build_daishi_with_no_courses_still_makes_a_sheet():
    wb = _wb([])
    assert wb.sheetnames == ["実績表"]
    assert wb["実績表"]["A1"].value == "6/26 ／ 1号　京阪北版"


def test_build_daishi_wrapping_at_per_row_boundary():
    """🔴 2026-09-04: 1コース=1列。per_row を超えたら次の3行ブロックへ折り返す。"""
    courses = [
        {"code": f"01010{i}", "chiku_name": "枚方・交野",
         "course_name": f"01010{i} 枚方・交野",
         "flyers": [{"name": "A", "count": 100}]}
        for i in range(1, 6)
    ]
    data = A.build_jisseki_daishi_workbook(courses, A.KEIHAN_KITA, gou=1,
                                           haifubi="2026-06-26", per_row=4)
    import io
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["実績表"]
    # 4コース目 grp0 col3 → D列, ヘッダー行2
    assert ws["D2"].value == "4\n010104"
    # 5コース目 grp1 col0 → A列, ヘッダー行 = 2 + 1*3 = 5
    assert ws["A5"].value == "5\n010105"


def test_build_daishi_default_per_row_is_20():
    """🔴 2026-09-04大橋様ご依頼: 実物どおり既定20列で折り返す(A3印刷に収める)。"""
    courses = [
        {"code": f"{i:06d}", "chiku_name": "枚方・交野",
         "course_name": f"{i:06d} 枚方・交野",
         "flyers": [{"name": "A", "count": 100}]}
        for i in range(1, 22)
    ]
    data = A.build_jisseki_daishi_workbook(courses, A.KEIHAN_KITA, gou=1,
                                           haifubi="2026-06-26")
    import io
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["実績表"]
    assert ws["T2"].value == "20\n000020"    # 20コース目=T列(20番目の列)、まだ1ブロック目
    assert ws["A5"].value == "21\n000021"    # 21コース目で次のブロックへ折り返す
