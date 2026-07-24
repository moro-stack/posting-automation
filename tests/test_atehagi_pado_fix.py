"""関西ぱど指摘反映（2026-07-23）の受け入れテスト。

設計書: docs/superpowers/specs/2026-07-23-atehagi-kansai-pado-fix-design.md
"""
import io

import openpyxl

from common import atehagi as A


def _kita_groups():
    from tests.test_atehagi import _sample_table
    return A.group_by_chiku(A.rows_from_table(_sample_table()))


def _minami_table():
    """南版（1211号）を模した合成テーブル。列構成は北版と同一。"""
    header = ["配布日", "号数", "ルート", "異動", "配送順位", "ぱどんな", "住所",
              "電話番号", "担当地区", "チラシコード", "配送物", "配布部数", "配送備考",
              "町界名", "街区（番地）名称", "受注種別", "チラシサイズ"]
    return [
        ["配送管理表", "2026-06-19", "(1211号)"],
        header,
        ["2026-06-19", 1211, 0, None, None, "ケイピーエス", None, None, 911001, None,
         "91 ぱど", 180, " ", "橋波西之町", "1丁目", None, None],
        ["2026-06-19", 1211, 0, None, None, "ケイピーエス", None, None, 911001, 31,
         "サンプルチラシ", 180, " ", "橋波西之町", "1丁目", "全戸配布(チラシ)", "Ａ４"],
        ["2026-06-19", 1211, 0, None, None, "フィールドサービス", None, None, 921203, None,
         "92 ぱど", 240, " ", "幸町", "2丁目", None, None],
    ]


def _minami_groups():
    return A.group_by_chiku(A.rows_from_table(_minami_table()))


def test_chiku_name_minami_is_moriguchi_kadoma():
    """南版の見出し地区名は担当地区コードによらず常に「守口・門真」。"""
    assert A.chiku_name(A.KEIHAN_MINAMI, "911001") == "守口・門真"
    assert A.chiku_name(A.KEIHAN_MINAMI, 921203) == "守口・門真"


def test_area_code_keeps_six_digit_minami_code():
    """南版の担当地区は6桁。従来の「末尾5桁」正規化では先頭の9が落ちていた。"""
    assert A.area_code("911001") == "911001"
    assert A.area_code(921203) == "921203"


def test_area_code_keeps_kita_five_digit_code_as_is():
    """北版5桁は従来どおり（シート名・グループ化キーを変えないため）。"""
    assert A.area_code(10101) == "10101"
    assert A.area_code("  46501 ") == "46501"
    assert A.area_code("X-44408") == "44408"      # 数字以外は除去


def test_area_code6_zero_pads_to_six_digits():
    """あて紙に印字するコードは6桁ゼロ埋め。先頭2桁が配送物名の番号と一致する。"""
    assert A.area_code6(10101) == "010101"        # → 配送物「01 ぱど」
    assert A.area_code6("46501") == "046501"      # → 配送物「04 ぱど」
    assert A.area_code6("911001") == "911001"     # 南版は既に6桁なので0は足さない
    assert A.area_code6(921203) == "921203"


def test_atehagi_prints_six_digit_chiku_and_clears_fixed_zero():
    """担当地区は6桁ゼロ埋めの1つのセルに印字し、テンプレ固定の「0」は出さない。"""
    data = A.build_atehagi_workbook(_kita_groups(), A.KEIHAN_KITA)
    ws = openpyxl.load_workbook(io.BytesIO(data))["10101"]
    assert ws["D4"].value == "010101"
    assert ws["C4"].value in (None, "")          # 固定の0は消す（二重表示の防止）
    assert str(ws["A4"].value).startswith("担当地区")   # ラベルは残す


def test_minami_atehagi_keeps_six_digit_code_and_fixed_area_name():
    """南版は先頭の9を落とさず、見出しは「守口・門真」。"""
    groups = _minami_groups()
    assert list(groups.keys()) == ["911001", "921203"]     # 従来は 11001/21203 になっていた
    data = A.build_atehagi_workbook(groups, A.KEIHAN_MINAMI)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["911001", "921203"]
    ws = wb["911001"]
    assert ws["A1"].value == "守口・門真"
    assert ws["D4"].value == "911001"            # 6桁のまま・0は足さない
    assert ws["C4"].value in (None, "")


def _shuffled_table():
    """ぱど行が先頭に無い配送管理表（並び順が変わっても正しく拾えるかの検証用）。"""
    header = ["配布日", "号数", "ぱどんな", "担当地区", "配送物", "配布部数", "チラシサイズ"]
    return [
        ["配送管理表", "2026-06-26", "(1248号)"],
        header,
        # チラシ行が先・部数も名前もぱど行と異なる
        ["2026-06-26", 1248, "チラシ担当", 10101, "サンプル生協", 999, "Ｂ４"],
        ["2026-06-26", 1248, "ぱど担当", 10101, "01 ぱど", 224, None],
        ["2026-06-26", 1248, "チラシ担当", 10101, "別のチラシ", 888, "Ａ４"],
    ]


def test_pado_row_is_picked_regardless_of_order():
    rows = A.rows_from_table(_shuffled_table())
    assert A.pado_row(rows)["haisoubutsu"] == "01 ぱど"
    assert A.pado_row(rows)["busuu"] == 224


def test_pado_row_falls_back_to_first_row_when_absent():
    rows = A.rows_from_table([
        ["担当地区", "配送物", "配布部数"],
        [10101, "チラシのみ", 300],
    ])
    assert A.pado_row(rows)["haisoubutsu"] == "チラシのみ"


def test_atehagi_header_uses_pado_row_not_first_row():
    """上部の部数・リーダー名は「ぱど行」の値。並び順に依存しない。"""
    groups = A.group_by_chiku(A.rows_from_table(_shuffled_table()))
    data = A.build_atehagi_workbook(groups, A.KEIHAN_KITA)
    ws = openpyxl.load_workbook(io.BytesIO(data))["10101"]
    assert ws["G2"].value == 224          # ぱどの部数（先頭行の999ではない）
    assert ws["A3"].value == "ぱど担当"    # ぱど行の名前
    # 明細の部数は各行のL列をそのまま（＝上部とは別の数字）
    assert [ws[f"J{r}"].value for r in (5, 6, 7)] == [999, 224, 888]


def test_jisseki_uses_pado_row_for_leader_and_busuu():
    """実績表のリーダー・部数もぱど行基準（あて紙と食い違わないように）。"""
    groups = A.group_by_chiku(A.rows_from_table(_shuffled_table()))
    rows = A.jisseki_rows(groups, A.KEIHAN_KITA)
    assert len(rows) == 1
    assert rows[0]["リーダー"] == "ぱど担当"
    assert rows[0]["部数"] == 224
    assert rows[0]["担当地区"] == "010101"       # あて紙と同じ6桁ゼロ埋め


def test_shukei_uses_pado_row_for_leader_and_busuu():
    """集計表のリーダー・部数もぱど行基準。"""
    groups = A.group_by_chiku(A.rows_from_table(_shuffled_table()))
    data = A.shukei_data(groups, A.KEIHAN_KITA)
    assert data["total_busuu"] == 224            # ぱどの部数（先頭行の999ではない）
    assert list(data["per"].keys()) == [("1", "ぱど担当")]


def test_atehagi_is_centered_on_paper():
    """B5用紙の左右・上下ともに中央に印刷されること（関西ぱど 2026-07-23）。"""
    data = A.build_atehagi_workbook(_kita_groups(), A.KEIHAN_KITA)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for name in wb.sheetnames:
        ws = wb[name]
        assert ws.print_options.horizontalCentered is True, f"{name}: 左右中央でない"
        assert ws.print_options.verticalCentered is True, f"{name}: 上下中央でない"
        assert ws.page_setup.paperSize == 13, f"{name}: B5(JIS)でない"


def test_parse_area_codes_normalizes_and_dedupes():
    """入力はカンマ/空白/改行区切り。ゼロ埋めの有無を問わず6桁に正規化する。"""
    assert A.parse_area_codes("10101") == ["010101"]
    assert A.parse_area_codes("010101") == ["010101"]          # 0付きでも同じ
    assert A.parse_area_codes("10101, 46501") == ["010101", "046501"]
    assert A.parse_area_codes("10101\n46501　911001") == ["010101", "046501", "911001"]
    assert A.parse_area_codes("10101,10101") == ["010101"]     # 重複は除く
    assert A.parse_area_codes("") == []


def test_select_groups_picks_only_requested_areas():
    groups = _kita_groups()                     # 10101 と 46501
    sel = A.select_groups(groups, "46501")
    assert list(sel.keys()) == ["46501"]
    sel2 = A.select_groups(groups, "046501, 010101")
    assert list(sel2.keys()) == ["46501", "10101"]    # 指定した順で返す
    assert A.select_groups(groups, "99999") == {}     # 無い地区は空


def test_single_area_workbook_has_one_sheet():
    groups = _kita_groups()
    sel = A.select_groups(groups, "10101")
    data = A.build_atehagi_workbook(sel, A.KEIHAN_KITA)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["10101"]
    assert wb["10101"]["D4"].value == "010101"


def test_atehagi_filename_includes_area_code():
    from tests.test_atehagi import _sample_table
    rows = A.rows_from_table(_sample_table())
    assert A.atehagi_filename(A.KEIHAN_KITA, rows, chiku="10101") == \
        "京阪北版_あて紙_1248号_20260626_010101.xlsx"
    assert A.atehagi_filename(A.KEIHAN_KITA, rows, chiku=["10101", "46501"]) == \
        "京阪北版_あて紙_1248号_20260626_010101他1件.xlsx"
    # 従来どおり（全地区）の呼び出しは変わらない
    assert A.atehagi_filename(A.KEIHAN_KITA, rows) == "京阪北版_あて紙_1248号_20260626.xlsx"


def _overflow_table(n_chirashi):
    """1地区に大量のチラシがある配送管理表（テンプレ13行の上限検証用）。"""
    header = ["配布日", "号数", "ぱどんな", "担当地区", "配送物", "配布部数", "チラシサイズ"]
    rows = [["配送管理表", "2026-06-26", "(1248号)"], header,
            ["2026-06-26", 1248, "太郎", 10101, "01 ぱど", 224, None]]
    for i in range(n_chirashi):
        rows.append(["2026-06-26", 1248, "太郎", 10101, f"チラシ{i + 1}", 224, "Ａ４"])
    return rows


def test_overflow_areas_detects_areas_beyond_template_rows():
    """テンプレは明細13行まで。それを超える地区を検出できること（無警告の欠落を防ぐ）。"""
    ok = A.group_by_chiku(A.rows_from_table(_overflow_table(12)))     # ぱど1+チラシ12=13行
    assert A.overflow_areas(ok) == []
    ng = A.group_by_chiku(A.rows_from_table(_overflow_table(15)))     # 合計16行
    assert A.overflow_areas(ng) == [("10101", 16)]


# ---------------------------------------------------------------------------
# 敵対的レビュー（2026-07-23）で実データ再現が確認された指摘への回帰テスト
# ---------------------------------------------------------------------------

def test_area_code_accepts_full_width_digits():
    """日本語入力のまま全角で打っても同じコードとして扱う（刷り直しの実運用で必須）。"""
    assert A.area_code("０１０１０１") == "010101"
    assert A.area_code6("１０１０１") == "010101"
    assert A.parse_area_codes("０１０１０１") == ["010101"]
    assert A.parse_area_codes("１０１０１，４６５０１") == ["010101", "046501"]


def test_parse_area_codes_rejects_too_long_token():
    """カンマ打ち忘れ（010101046501）で片方が黙って消えないこと。"""
    assert A.parse_area_codes("010101046501") == []
    assert A.parse_area_codes("010101, 046501") == ["010101", "046501"]


def test_chiku_name_kita_handles_zero_padded_text_code():
    """先頭0付きの6桁テキストで来ても北版の地区名が空にならないこと。"""
    assert A.chiku_name(A.KEIHAN_KITA, "010101") == "枚方・交野"
    assert A.chiku_name(A.KEIHAN_KITA, "044408") == "寝屋川・枚方"
    assert A.chiku_name(A.KEIHAN_KITA, "10101") == "枚方・交野"      # 従来形も維持


def test_shukei_area_does_not_collapse_for_minami():
    """南版の集計表でエリア別が1行に潰れないこと（従来は全件"9"で1行になっていた）。

    ラベルの値そのものは test_shukei_area_keeps_csv_numbers_for_minami で固定している。
    """
    table = [
        ["担当地区", "ぱどんな", "配送物", "配布部数", "チラシサイズ"],
        [911001, "ケイピーエス", "91 ぱど", 100, ""],
        [921203, "ケイピーエス", "92 ぱど", 200, ""],
    ]
    groups = A.group_by_chiku(A.rows_from_table(table))
    data = A.shukei_data(groups, A.KEIHAN_MINAMI)
    assert len(data["area_busuu"]) == 2, data["area_busuu"]


def test_shukei_area_unchanged_for_kita():
    """北版の集計表のエリア区分は従来どおり（1〜5）。"""
    groups = _kita_groups()          # 10101 と 46501
    data = A.shukei_data(groups, A.KEIHAN_KITA)
    assert sorted(data["area_busuu"].keys()) == ["1", "4"]


def test_detect_version_from_area_codes():
    """担当地区コードから版を判定できること（版の取り違えを止めるため）。"""
    assert A.detect_version(_kita_groups()) == A.KEIHAN_KITA
    assert A.detect_version(_minami_groups()) == A.KEIHAN_MINAMI
    assert A.detect_version(A.group_by_chiku(A.rows_from_table([
        ["担当地区", "配送物", "配布部数"], [777777, "謎", 1]]))) is None


def test_details_are_truncated_to_template_rows_without_overrunning():
    """明細がテンプレ上限を超えたとき、集計行(18行目)を侵食しないこと。

    切り捨て自体はテンプレの枠が13行しかないため避けられない（UIで警告する）。
    ここでは「上限まで埋まる」「上限を1行でも超えて書かない」ことを固定する。
    """
    groups = A.group_by_chiku(A.rows_from_table(_overflow_table(15)))   # ぱど1+チラシ15=16行
    data = A.build_atehagi_workbook(groups, A.KEIHAN_KITA)
    ws = openpyxl.load_workbook(io.BytesIO(data))["10101"]
    last = 5 + A._MAX_DETAIL_ROWS - 1
    assert ws[f"B{last}"].value is not None, "上限行まで埋まっていない"
    assert ws[f"B{last + 1}"].value is None, "上限を超えて書き込んでいる（集計行を侵食）"
    assert ws["I18"].value == 15, "チラシ数は切り捨て前の全件を数える（警告と整合させるため）"


def test_area_no_uses_csv_value_as_is():
    """エリアは担当地区コードの値をそのまま使う。勝手に振り直さない（オーナー指示 2026-07-23）。

    北版 10101 → エリア「1」（従来どおり）／南版 911001 → エリア「91」（「1」に読み替えない）。
    先頭2桁は配送物名の番号（「01 ぱど」「91 ぱど」）と一致する。
    """
    assert A.area_no(10101) == "1"
    assert A.area_no("46501") == "4"
    assert A.area_no("911001") == "91"
    assert A.area_no(921203) == "92"
    assert A.area_no("010101") == "1"      # 先頭0付きテキストでも同じ


def test_shukei_area_keeps_csv_numbers_for_minami():
    """南版の集計表のエリアは 91／92（1／2に振り直さない）。"""
    table = [
        ["担当地区", "ぱどんな", "配送物", "配布部数", "チラシサイズ"],
        [911001, "ケイピーエス", "91 ぱど", 100, ""],
        [921203, "ケイピーエス", "92 ぱど", 200, ""],
    ]
    groups = A.group_by_chiku(A.rows_from_table(table))
    data = A.shukei_data(groups, A.KEIHAN_MINAMI)
    assert sorted(data["area_busuu"].items()) == [("91", 100), ("92", 200)]
