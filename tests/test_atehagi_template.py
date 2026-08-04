import os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(__file__))
TMPL = os.path.join(ROOT, "templates", "あて紙テンプレート_京阪.xlsx")


def test_template_has_constants_merges_and_no_data():
    wb = openpyxl.load_workbook(TMPL)
    ws = wb[wb.sheetnames[0]]
    merged = {str(m) for m in ws.merged_cells.ranges}
    assert {"A3:E3", "D4:G4", "G2:I2"} <= merged
    assert str(ws["A4"].value).startswith("担当地区")   # 固定ラベル保持
    assert ws["C4"].value == 0
    assert str(ws["J2"].value).strip() == "部"
    # 可変セルは空（実データが残っていない）
    for coord in ["A1", "A3", "D4", "G2", "B5", "B6", "I6", "J5", "I18"]:
        assert ws[coord].value in (None, ""), f"{coord} にデータが残っている"
