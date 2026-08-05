import hashlib
import io
import time

import openpyxl
import pytest
from common import invoice_excel


def _load(data):
    return openpyxl.load_workbook(io.BytesIO(data)).active


def test_build_invoice_fills_lines_and_total():
    lines = [
        {"project_name": "関西ぱど：京阪北版", "report_qty": 3713, "unit_price": 3,
         "amount": 11139, "remark": "配布"},
        {"project_name": "関西ぱど：京阪北版", "report_qty": 1, "unit_price": 540,
         "amount": 540, "remark": "交通費"},
    ]
    data = invoice_excel.build_invoice_xlsx(
        distributor_name="時野", issue_date="2026年6月30日",
        period_from="2026年6月19日", period_to="2026年6月27日", lines=lines)
    ws = _load(data)
    # 明細1行目(13行) の案件名・数量・単価・計
    assert ws["B13"].value == "関西ぱど：京阪北版"
    assert ws["D13"].value == 3713
    assert ws["E13"].value == 3
    assert ws["F13"].value == 11139
    # ご請求金額(A7)に合計 11679 が含まれる(カンマ区切り表記を許容)
    assert "11679" in str(ws["A7"].value).replace(",", "")
    # 配布部数(F20)= 配布のみ 3713(交通費は除外、単位「部」・カンマ区切りを許容)
    assert "3713" in str(ws["F20"].value).replace(",", "")


def test_build_invoice_fills_remark_into_bikou_column():
    """種別(配布/交通費など)を備考(G列)へ各行に入れる。"""
    lines = [
        {"project_name": "関西ぱど：京阪北版", "report_qty": 3713, "unit_price": 3,
         "amount": 11139, "remark": "配布"},
        {"project_name": "関西ぱど：京阪北版", "report_qty": 1, "unit_price": 540,
         "amount": 540, "remark": "交通費"},
    ]
    data = invoice_excel.build_invoice_xlsx(
        distributor_name="時野", issue_date="2026年6月30日",
        period_from="2026年6月19日", period_to="2026年6月27日", lines=lines)
    ws = _load(data)
    assert ws["G12"].value == "備考"      # 見出しはテンプレート由来
    assert ws["G13"].value == "配布"
    assert ws["G14"].value == "交通費"


def test_build_invoice_leaves_unused_remark_rows_empty():
    """明細が1行だけなら、残りの備考行にテンプレートの例文が残らない。"""
    data = invoice_excel.build_invoice_xlsx(
        distributor_name="時野", issue_date="2026年6月30日",
        period_from="", period_to="",
        lines=[{"project_name": "その他", "report_qty": 10, "unit_price": 5,
                "amount": 50, "remark": "配布"}])
    ws = _load(data)
    assert ws["G13"].value == "配布"
    for r in range(14, 19):
        assert ws[f"G{r}"].value in (None, "")


def test_build_invoice_writes_isshiki_in_qty_column_for_non_delivery():
    """請求書の数量欄(D列)は、交通費・手当・その他なら数字ではなく『一式』と書く。
    配布・挟み込みは部数を数えるので数値のまま(数量×単価=計 が読めるように)。"""
    lines = [
        {"project_name": "関西ぱど：京阪北版", "report_qty": 3713, "unit_price": 3,
         "amount": 11139, "remark": "配布"},
        {"project_name": "関西ぱど：京阪北版", "report_qty": 1, "unit_price": 540,
         "amount": 540, "remark": "交通費"},
        {"project_name": "関西ぱど：京阪北版", "report_qty": 1, "unit_price": 1000,
         "amount": 1000, "remark": "手当"},
    ]
    data = invoice_excel.build_invoice_xlsx(
        distributor_name="時野", issue_date="2026年6月30日",
        period_from="2026年6月19日", period_to="2026年6月27日", lines=lines)
    ws = _load(data)
    assert ws["D13"].value == 3713      # 配布は数値のまま
    assert ws["D14"].value == "一式"     # 交通費
    assert ws["D15"].value == "一式"     # 手当
    # 単価・計は従来どおり数値
    assert ws["E14"].value == 540
    assert ws["F14"].value == 540


def test_build_invoice_is_deterministic():
    """同じ内容なら毎回まったく同じバイト列を返すこと。

    Streamlit のダウンロードURLはファイル内容のハッシュで決まる。生成のたびに
    バイト列が変わると、再実行のたびに別URLが発行され、ブラウザが取得中の
    古いURLがサーバー側から破棄されて 404 になる（＝Excelが落ちてこない）。
    """
    kw = dict(distributor_name="時野", issue_date="2026-07-15",
              period_from="2026-07-15", period_to="2026-07-15",
              lines=[{"project_name": "関西ぱど：京阪北版", "report_qty": 100,
                      "unit_price": 4, "amount": 400, "remark": "挟み込み"}])
    first = invoice_excel.build_invoice_xlsx(**kw)
    time.sleep(1.1)   # ZIP内のタイムスタンプは秒単位。1秒以上あけて差が出ないか見る
    second = invoice_excel.build_invoice_xlsx(**kw)
    assert hashlib.md5(first).hexdigest() == hashlib.md5(second).hexdigest()


def test_build_invoice_stays_a_valid_xlsx():
    """バイト列を安定させても、Excelとして正しく開けて中身が読めること。"""
    data = invoice_excel.build_invoice_xlsx(
        distributor_name="時野", issue_date="2026-07-15",
        period_from="2026-07-15", period_to="2026-07-15",
        lines=[{"project_name": "関西ぱど：京阪北版", "report_qty": 100,
                "unit_price": 4, "amount": 400, "remark": "挟み込み"}])
    ws = _load(data)
    assert ws["B13"].value == "関西ぱど：京阪北版"
    assert ws["D13"].value == 100
    assert ws["F13"].value == 400
    assert ws["G13"].value == "挟み込み"
    assert ws["A1"].value.startswith("業務完了報告書")     # テンプレの書式が生きている
    assert ws["G12"].value == "備考"


def test_build_invoice_returns_bytes():
    data = invoice_excel.build_invoice_xlsx(
        distributor_name="黒瀬", issue_date="2026年6月30日",
        period_from="", period_to="",
        lines=[{"project_name": "その他", "report_qty": 10, "unit_price": 5,
                "amount": 50, "remark": "配布"}])
    assert isinstance(data, (bytes, bytearray)) and len(data) > 0


def test_build_invoice_rejects_more_than_max_line_rows():
    lines = [{"project_name": "その他", "report_qty": 1, "unit_price": 1,
              "amount": 1, "remark": "配布"} for _ in range(7)]
    with pytest.raises(ValueError):
        invoice_excel.build_invoice_xlsx(
            distributor_name="黒瀬", issue_date="2026年6月30日",
            period_from="", period_to="", lines=lines)


def test_invoice_xlsx_nichito_puts_days_in_remark():
    """日当の請求書は備考に「配布（3 日）」と単位が出る。数量セルは日数の数値。"""
    xlsx = invoice_excel.build_invoice_xlsx(
        distributor_name="山田太郎", issue_date="2026-07-17",
        period_from="2026-07-01", period_to="2026-07-15",
        lines=[{"project_name": "案件A", "report_qty": 3, "unit_price": 8000,
                "amount": 24000, "remark": "配布", "copies": 3713}],
        pay_type="日当")
    ws = _load(xlsx)
    assert ws["D13"].value == 3
    assert ws["G13"].value == "配布（3 日）"
    # 報告数は copies 列から取る(数量の3ではない)
    assert ws["F20"].value == "3,713部"


def test_invoice_xlsx_jikyu_puts_hours_in_remark():
    """時給は備考に「配布（8 時間）」と出る。報告数は copies から。"""
    xlsx = invoice_excel.build_invoice_xlsx(
        distributor_name="時給の人", issue_date="2026-07-17",
        period_from="2026-07-01", period_to="2026-07-15",
        lines=[{"project_name": "案件A", "report_qty": 8, "unit_price": 1200,
                "amount": 9600, "remark": "配布", "copies": 1200}],
        pay_type="時給")
    ws = _load(xlsx)
    assert ws["D13"].value == 8
    assert ws["G13"].value == "配布（8 時間）"
    assert ws["F20"].value == "1,200部"


def test_invoice_xlsx_getkyu_shows_isshiki():
    """月給は数量を数えないので数量セルが「一式」。"""
    xlsx = invoice_excel.build_invoice_xlsx(
        distributor_name="月給の人", issue_date="2026-07-17",
        period_from="2026-07-01", period_to="2026-07-31",
        lines=[{"project_name": "案件A", "report_qty": 1, "unit_price": 250000,
                "amount": 250000, "remark": "配布", "copies": 5000}],
        pay_type="月給")
    ws = _load(xlsx)
    assert ws["D13"].value == "一式"
    assert ws["G13"].value == "配布"       # 一式の行に単位は併記しない
    assert ws["F20"].value == "5,000部"


def test_invoice_xlsx_nichito_does_not_add_unit_to_non_delivery_rows():
    """日当でも、交通費・手当は数を数えない(一式)ので備考に単位を併記しない。"""
    xlsx = invoice_excel.build_invoice_xlsx(
        distributor_name="山田太郎", issue_date="2026-07-17",
        period_from="2026-07-01", period_to="2026-07-15",
        lines=[{"project_name": "案件A", "report_qty": 3, "unit_price": 8000,
                "amount": 24000, "remark": "配布", "copies": 3713},
               {"project_name": "案件A", "report_qty": 1, "unit_price": 540,
                "amount": 540, "remark": "交通費"}],
        pay_type="日当")
    ws = _load(xlsx)
    assert ws["G13"].value == "配布（3 日）"
    assert ws["D14"].value == "一式"
    assert ws["G14"].value == "交通費"


def test_invoice_xlsx_without_pay_type_keeps_current_behavior():
    """pay_type を渡さない既存の呼び出しは1ミリも変わらない(後方互換)。"""
    xlsx = invoice_excel.build_invoice_xlsx(
        distributor_name="山田太郎", issue_date="2026-07-17",
        period_from="2026-07-01", period_to="2026-07-15",
        lines=[{"project_name": "案件A", "report_qty": 3713, "unit_price": 2.5,
                "amount": 9283, "remark": "配布"},
               {"project_name": "案件A", "report_qty": 1, "unit_price": 540,
                "amount": 540, "remark": "交通費"}])
    ws = _load(xlsx)
    assert ws["D13"].value == 3713          # 配布は数値のまま
    assert ws["G13"].value == "配布"         # 備考に単位は付けない
    assert ws["D14"].value == "一式"         # 交通費は今まで通り
    assert ws["F20"].value == "3,713部"


def test_invoice_xlsx_houbai_pay_type_matches_no_pay_type():
    """歩合を明示しても、pay_type なしと同じ出力になること。"""
    kw = dict(distributor_name="山田太郎", issue_date="2026-07-17",
              period_from="2026-07-01", period_to="2026-07-15",
              lines=[{"project_name": "案件A", "report_qty": 3713, "unit_price": 2.5,
                      "amount": 9283, "remark": "配布"}])
    assert (hashlib.md5(invoice_excel.build_invoice_xlsx(**kw, pay_type="歩合")).hexdigest()
            == hashlib.md5(invoice_excel.build_invoice_xlsx(**kw)).hexdigest())


def test_freeze_xlsx_bytes_stabilises_pandas_output():
    """一覧の「Excelで保存」(pandas出力)も、同じ内容なら同じバイト列になること。"""
    import pandas as pd
    from common.excel_io import freeze_xlsx_bytes

    def build():
        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as w:
            pd.DataFrame([{"日付": "2026-07-15", "金額": "¥400"}]).to_excel(
                w, index=False, sheet_name="data")
        return freeze_xlsx_bytes(buf.getvalue())

    first = build()
    time.sleep(1.1)
    second = build()
    assert hashlib.md5(first).hexdigest() == hashlib.md5(second).hexdigest()
    # 固定しても中身は読める
    ws = _load(first)
    assert ws["A1"].value == "日付"
    assert ws["B2"].value == "¥400"


def test_build_invoice_print_area_includes_bikou_column():
    """🔴 依頼①。備考(G列)が印刷範囲に入っていること。

    テンプレートの print_area は 'Sheet1'!$A$1:$F$27 で、G列が印刷範囲の外にあった。
    そのためA4印刷すると備考だけ紙に載らない(幅不足ではない)。
    """
    data = invoice_excel.build_invoice_xlsx(
        distributor_name="時野", issue_date="2026-08-05",
        period_from="2026-08-01", period_to="2026-08-04",
        lines=[{"project_name": "案件A", "report_qty": 100, "unit_price": 4,
                "amount": 400, "remark": "配布"}])
    ws = _load(data)
    assert ws.print_area == "'Sheet1'!$A$1:$G$27"


def test_build_invoice_fits_to_one_page_wide():
    """横は必ず1ページに収める。列を増やしても切れないようにするため。"""
    data = invoice_excel.build_invoice_xlsx(
        distributor_name="時野", issue_date="2026-08-05",
        period_from="", period_to="",
        lines=[{"project_name": "案件A", "report_qty": 100, "unit_price": 4,
                "amount": 400, "remark": "配布"}])
    ws = _load(data)
    assert ws.page_setup.fitToWidth == 1
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True


def test_build_invoice_all_written_bikou_rows_are_inside_print_area():
    """🔴 列だけでなく行も見る。備考を書いた明細行が全部 print_area の内側にあること。
    明細は13〜18行なので、print_area の下端が12行などに退行したら落ちる。"""
    lines = [{"project_name": f"案件{i}", "report_qty": 1, "unit_price": 1,
              "amount": 1, "remark": "配布"} for i in range(6)]
    data = invoice_excel.build_invoice_xlsx(
        distributor_name="時野", issue_date="2026-08-05",
        period_from="", period_to="", lines=lines)
    ws = _load(data)

    from openpyxl.utils.cell import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(
        ws.print_area.split("!")[-1].replace("$", ""))

    written = [(c, r) for r in range(1, ws.max_row + 1)
               for c in range(1, ws.max_column + 1)
               if ws.cell(r, c).value not in (None, "", " ", "\u3000")]
    assert written, "セルに何も書かれていない＝テストが空振りしている"
    outside = [(c, r) for c, r in written
               if not (min_col <= c <= max_col and min_row <= r <= max_row)]
    assert outside == [], f"印刷範囲の外に中身がある: {outside}"
