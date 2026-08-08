"""仕分け表（倉庫でチラシを配布員ごとの山に分けるチェック表）のテスト。

設計: docs/superpowers/specs/2026-08-07-shiwake-hyo-design.md
"""
import datetime as dt
import io
import os
from pathlib import Path

import openpyxl
import pytest

from common import shiwake

HEADER = ["配布日", "号数", "ルート", "異動", "配送順位", "ぱどんな", "住所", "電話番号",
          "担当地区", "チラシコード", "配送物", "配布部数", "配送備考", "町界名",
          "街区（番地）名称", "受注種別", "チラシサイズ"]


def _row(*, padonna, chirashi, busuu, size=None, junni=None, ido=None, chiku="911103"):
    r = [None] * len(HEADER)
    r[0] = dt.datetime(2026, 8, 21)
    r[1] = "1213"
    r[3] = ido
    r[4] = junni
    r[5] = padonna
    r[8] = chiku
    r[10] = chirashi
    r[11] = busuu
    r[16] = size
    return r


def _table(rows):
    return [["配送管理表", None], HEADER] + rows


# ===== 1. 配布員ごとに合算する =====


def test_busuu_is_summed_per_person_across_chiku():
    """🔴 松岡さんは2地区持ち。業務スーパーは 438+693=1131 と1行にまとまること。"""
    table = _table([
        _row(padonna="松岡  直美", chirashi="91 ぱど", busuu=438, chiku="911103",
             junni=dt.datetime(9101, 4, 1)),
        _row(padonna="松岡  直美", chirashi="業務スーパー", busuu=438, size="Ｂ４",
             chiku="911103", junni=dt.datetime(9101, 4, 1)),
        _row(padonna="松岡  直美", chirashi="91 ぱど", busuu=693, chiku="911107",
             junni=dt.datetime(9101, 4, 1)),
        _row(padonna="松岡  直美", chirashi="業務スーパー", busuu=693, size="Ｂ４",
             chiku="911107", junni=dt.datetime(9101, 4, 1)),
    ])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    assert len(groups) == 1
    items = {i["chirashi"]: i["busuu"] for i in groups[0]["items"]}
    assert items == {"91 ぱど": 1131, "業務スーパー": 1131}


def test_group_total_matches_sum_of_items():
    table = _table([
        _row(padonna="A", chirashi="X", busuu=100),
        _row(padonna="A", chirashi="Y", busuu=250, size="Ｂ４"),
    ])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    assert groups[0]["total"] == 350
    assert groups[0]["total"] == sum(i["busuu"] for i in groups[0]["items"])


# ===== 2〜3. 休の扱い =====


def test_absent_person_is_excluded():
    """🔴 異動＝休 の人は仕分け表に入れない(オーナー指示)。"""
    table = _table([
        _row(padonna="働く人", chirashi="X", busuu=10),
        _row(padonna="休む人", chirashi="X", busuu=10, ido="休"),
    ])
    groups, warn = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    names = [g["name"] for g in groups]
    assert names == ["働く人"]
    assert warn["excluded"] == ["休む人"]


def test_unknown_ido_value_is_kept_and_warned():
    """🔴 見慣れない異動値で黙って人を落とさない。残したうえで警告する。

    勝手に除外して人数が減るのが一番こわい(倉庫で山が足りなくなる)。"""
    table = _table([
        _row(padonna="働く人", chirashi="X", busuu=10),
        _row(padonna="謎の人", chirashi="X", busuu=10, ido="移"),
    ])
    groups, warn = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    names = sorted(g["name"] for g in groups)
    assert names == ["働く人", "謎の人"]
    assert warn["unknown_ido"] == [("謎の人", "移")]


# ===== 4〜5. チラシサイズ =====


def test_empty_size_becomes_johoshi():
    """🔴 チラシサイズが空＝ぱど本誌 → 「情報誌」と書く。"""
    table = _table([_row(padonna="A", chirashi="91 ぱど", busuu=100, size=None)])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    assert groups[0]["items"][0]["size"] == "情報誌"


@pytest.mark.parametrize("size", ["Ｂ３", "Ｂ４", "Ａ４"])
def test_existing_size_is_kept_as_is(size):
    """全角のまま。オーナー指示「エリア表記等は基本的にCSVをそのまま拾って」。"""
    table = _table([_row(padonna="A", chirashi="X", busuu=100, size=size)])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    assert groups[0]["items"][0]["size"] == size


# ===== 6. 配送順位の表記 =====


def test_junni_datetime_is_formatted_back_to_original():
    """🔴 Excelが 9101-04-01 を datetime(9101,4,1) として読む。
    そのまま出すと 9101/4/1 になって別物になるため元の表記に戻す。"""
    assert shiwake.format_junni(dt.datetime(9101, 4, 1)) == "9101-04-01"
    assert shiwake.format_junni(dt.datetime(9101, 12, 31)) == "9101-12-31"


def test_junni_string_is_kept_as_is():
    assert shiwake.format_junni("9101-04-01") == "9101-04-01"


def test_junni_none_becomes_empty():
    assert shiwake.format_junni(None) == ""


# ===== 7. 並び順 =====


def test_groups_are_sorted_by_junni_with_blanks_last():
    table = _table([
        _row(padonna="三番", chirashi="X", busuu=1, junni=dt.datetime(9101, 3, 1)),
        _row(padonna="空", chirashi="X", busuu=1, junni=None),
        _row(padonna="一番", chirashi="X", busuu=1, junni=dt.datetime(9101, 1, 1)),
        _row(padonna="二番", chirashi="X", busuu=1, junni=dt.datetime(9101, 2, 1)),
    ])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    assert [g["name"] for g in groups] == ["一番", "二番", "三番", "空"]


# ===== 8〜9. Excel 出力 =====


def _build(groups):
    data = shiwake.build_shiwake_workbook(groups, gou="1213", haifubi="2026-08-21")
    return openpyxl.load_workbook(io.BytesIO(data))


def test_workbook_has_one_sheet_per_person():
    """🔴 配布員ごとに1シート(2026-08-08 オーナー指示)。

    1人ずつ印刷する可能性があるため、Excelファイルは1つのままシートを分ける。
    """
    table = _table([
        _row(padonna="A", chirashi="X", busuu=1, junni=dt.datetime(9101, 1, 1)),
        _row(padonna="B", chirashi="X", busuu=1, junni=dt.datetime(9101, 2, 1)),
        _row(padonna="C", chirashi="X", busuu=1, junni=dt.datetime(9101, 3, 1)),
    ])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    wb = _build(groups)
    assert len(wb.sheetnames) == len(groups) == 3
    assert wb.sheetnames == ["A", "B", "C"]


def test_sheet_names_are_sanitized_and_unique():
    """🔴 シート名は31文字まで、Excelが禁じる記号は使えない。同名も付けられない。

    配布員名にこれらが入っていても落ちないこと。名前が消えないこと。
    """
    long_name = "あ" * 40
    table = _table([
        _row(padonna=long_name, chirashi="X", busuu=1, junni=dt.datetime(9101, 1, 1)),
        _row(padonna="メイト/南川 信哉", chirashi="X", busuu=1,
             junni=dt.datetime(9101, 2, 1)),
        _row(padonna="A[1]", chirashi="X", busuu=1, junni=dt.datetime(9101, 3, 1)),
    ])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    wb = _build(groups)
    names = wb.sheetnames
    assert len(names) == 3
    assert len(set(names)) == 3                     # 重複なし
    for n in names:
        assert len(n) <= 31
        assert not set(n) & set("[]:*?/" + chr(92))
    # スラッシュは消すのではなく別の字に置き換えて、誰の分か分かるようにする
    assert any("南川" in n for n in names)


def test_duplicate_person_names_get_distinct_sheets():
    """同姓同名がいてもシートが潰れないこと。"""
    table = _table([
        _row(padonna="田尻　美穂子", chirashi="X", busuu=1, chiku="911001",
             junni=dt.datetime(9101, 1, 1)),
    ])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    groups = groups + [dict(groups[0])]      # わざと同名を2件にする
    wb = _build(groups)
    assert len(wb.sheetnames) == 2
    assert len(set(wb.sheetnames)) == 2


def test_each_sheet_has_its_own_person_only():
    """🔴 シートに他人の情報が混ざらないこと。"""
    table = _table([
        _row(padonna="山田", chirashi="X", busuu=100, junni=dt.datetime(9101, 1, 1)),
        _row(padonna="鈴木", chirashi="Y", busuu=200, junni=dt.datetime(9101, 2, 1)),
    ])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    wb = _build(groups)
    for name in wb.sheetnames:
        text = chr(10).join(str(c.value) for row in wb[name].iter_rows()
                             for c in row if c.value is not None)
        others = [n for n in wb.sheetnames if n != name]
        assert name in text
        for o in others:
            assert o not in text


def test_workbook_shows_name_junni_and_items():
    table = _table([
        _row(padonna="松岡  直美", chirashi="91 ぱど", busuu=438,
             junni=dt.datetime(9101, 4, 1)),
        _row(padonna="松岡  直美", chirashi="91 ぱど", busuu=693,
             junni=dt.datetime(9101, 4, 1)),
    ])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    ws = _build(groups).active
    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row
                     if c.value is not None)
    assert "松岡  直美" in text
    assert "9101-04-01" in text
    assert "情報誌" in text
    assert "1131" in text or "1,131" in text


def test_workbook_check_column_is_empty():
    """チェック欄は空。作業員が手で書く。"""
    table = _table([_row(padonna="A", chirashi="X", busuu=1)])
    groups, _ = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    ws = _build(groups).worksheets[0]
    header_row = None
    for row in ws.iter_rows():
        vals = [c.value for c in row]
        if "チラシ名" in vals:
            header_row = row[0].row
            break
    assert header_row is not None
    check_col = [c.column for c in ws[header_row] if c.value == "済"][0]
    assert ws.cell(header_row + 1, check_col).value is None


# ===== 10. 実データ =====


REAL = r"C:\Users\moro\Downloads\【京阪南】20260821配送管理表 (1).xlsx"


@pytest.mark.skipif(not os.path.exists(REAL), reason="実データが無い環境ではスキップ")
def test_real_data_kyohan_minami_20260821():
    """🔴 実データ(京阪南 8/21号)で、休1名を除いた24名になること。"""
    wb = openpyxl.load_workbook(REAL, data_only=True)
    ws = wb["配送管理表"]
    table = [[c.value for c in row] for row in ws.iter_rows()]
    groups, warn = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))

    assert len(groups) == 24
    assert warn["excluded"] == ["井上\u3000美由紀"]
    assert warn["unknown_ido"] == []

    matsuoka = [g for g in groups if g["name"] == "松岡  直美"][0]
    items = {i["chirashi"]: (i["busuu"], i["size"]) for i in matsuoka["items"]}
    assert items["91 ぱど"] == (1131, "情報誌")
    assert items["業務スーパー"] == (1131, "Ｂ４")
    assert items["おたからや寝屋川店・萱島駅前店"] == (438, "Ｂ４")
    assert matsuoka["total"] == 1131 + 1131 + 438
    assert matsuoka["junni"] == "9101-04-01"


# ===== 11. 間違ったファイルを黙って通さない =====


def test_warns_when_no_row_has_junni():
    """🔴 「その他配送管理表」を入れると、配送順位が全行空になる。

    このファイルは配布員ではなく会社名(ケイピーエス/フィールドサービス)で
    まとまっているため、仕分け表として出すと「配布員2名・15万部」という
    無意味な表が黙って出来上がる。全行の配送順位が空なら知らせる。
    """
    table = _table([
        _row(padonna="ケイピーエス", chirashi="X", busuu=438, junni=None),
        _row(padonna="ケイピーエス", chirashi="Y", busuu=438, junni=None),
        _row(padonna="フィールドサービス", chirashi="X", busuu=100, junni=None),
    ])
    groups, warn = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    assert warn["looks_like_wrong_file"] is True
    # 中身は作る(判断はオーナーに委ねる)。黙って空にはしない。
    assert len(groups) == 2


def test_no_wrong_file_warning_when_junni_exists():
    table = _table([
        _row(padonna="A", chirashi="X", busuu=1, junni=dt.datetime(9101, 1, 1)),
        _row(padonna="B", chirashi="X", busuu=1, junni=None),
    ])
    _, warn = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    assert warn["looks_like_wrong_file"] is False


@pytest.mark.skipif(not os.path.exists(REAL), reason="実データが無い環境ではスキップ")
def test_real_haiso_kanri_hyo_is_not_flagged():
    """正しいファイル(配送管理表)では警告が出ないこと。"""
    wb = openpyxl.load_workbook(REAL, data_only=True)
    ws = wb["配送管理表"]
    table = [[c.value for c in row] for row in ws.iter_rows()]
    _, warn = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    assert warn["looks_like_wrong_file"] is False


SONOTA = r"C:\Users\moro\Downloads\京阪南_その他配送管理表 (1).CSV"


@pytest.mark.skipif(not os.path.exists(SONOTA), reason="実データが無い環境ではスキップ")
def test_real_sonota_haiso_is_flagged():
    """🔴 実物の「その他配送管理表」で警告が出ること。"""
    import csv as _csv
    text = Path(SONOTA).read_bytes().decode("cp932")
    table = list(_csv.reader(text.splitlines()))
    _, warn = shiwake.shiwake_groups(shiwake.rows_from_haiso_table(table))
    assert warn["looks_like_wrong_file"] is True
