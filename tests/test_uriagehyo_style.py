"""会議用売上表を、新しいテンプレート
`templates/KPS(大阪)原価売上表テンプレート.xlsx` のレイアウトどおりに作る。

依頼⑦(2026-08-27 大橋様)。これまでは「実物と同じ書式の1行1案件の表」を
アプリが独自に組み立てていたが、実際の会議用売上表は
  京阪南／京阪北 → 単独チラシ → アド・バリュー(週ごと) → リビング → 集計
というセクション構成になっている。テンプレートの構成に合わせて数字を流し込む。

目視で「合っている」と言わないための土台(集計表 test_shukei_style.py と同じ考え方)。
"""
import io

import openpyxl
import pytest

from common import uriagehyo as U
from common import uriagehyo_style as S


@pytest.fixture(scope="module")
def tpl():
    return S.UriagehyoTemplate()


# ===== テンプレートそのものの前提を固定する =====


def test_template_is_the_new_blank_one_from_ohashi_san(tpl):
    """新テンプレートは「1シートだけの白紙」。実データ入りの財務台帳
    (売上表テンプレート_大阪支社.xlsx・19シート)ではないこと。"""
    assert tpl.path.name == "KPS(大阪)原価売上表テンプレート.xlsx"
    assert tpl.ws.max_row >= 21


def test_template_has_the_shiwake_columns(tpl):
    """新テンプレートは「仕分け」列(G:H)を持つ(旧テンプレートの月によっては無い)。"""
    assert tpl.ws["G2"].value == "仕分け"
    assert tpl.ws["I2"].value == "その他"


def test_template_has_column_widths(tpl):
    widths = {k: d.width for k, d in tpl.ws.column_dimensions.items() if d.width}
    assert "A" in widths and "L" in widths


def test_template_marks_where_extra_single_flyer_rows_go(tpl):
    """単独チラシ行は「案件ごとに行増やす」前提だとテンプレート自身が書いている。"""
    assert "行増やす" in str(tpl.ws["A7"].value)


# ===== セクションの振り分け =====


def test_section_of_maps_project_names_to_sections():
    assert U.section_of("関西ぱど：京阪南版") == U.SECTION_PADO_MINAMI
    assert U.section_of("関西ぱど：京阪北版") == U.SECTION_PADO_KITA
    assert U.section_of("アドバリュー　8-1") == U.SECTION_ADVALUE
    assert U.section_of("リビングプロシード") == U.SECTION_LIVING
    assert U.section_of("その他　買取専科") == U.SECTION_TANDOKU


def test_section_of_puts_unknown_labels_in_the_unclassified_bucket():
    """🔴 案件マスタが増えたときに、どのセクションにも出ない行を作らない。"""
    assert U.section_of("知らない案件") == U.SECTION_UNKNOWN
    assert U.section_of("") == U.SECTION_UNKNOWN
    assert U.section_of(None) == U.SECTION_UNKNOWN


# ===== 生成した表 =====


def _row(uriage, genka, koutsuhi=0, nomimono=0):
    return U.build_row(hakko_gou=None, ban_mei=None, uriage_zeikomi=uriage,
                       genka_goukei_zeikomi=genka, koutsuhi=koutsuhi, nomimono=nomimono)


def _build(rows=None, year=2026, month=8):
    bulk = rows if rows is not None else [
        ("関西ぱど：京阪南版", _row(847502, 324505)),
        ("関西ぱど：京阪北版", _row(1804615, 466713)),
        ("アドバリュー　8-1", _row(544970, 177540, koutsuhi=6260)),
        ("アドバリュー　8-2", _row(214466, 87540, koutsuhi=400)),
        ("リビングプロシード", _row(173498, 90000)),
    ]
    data = S.build_monthly_workbook(bulk, year, month)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    return wb, wb.active


def _find_row(ws, col, value):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, col).value == value:
            return r
    return None


def test_sheet_title_matches_month():
    wb, ws = _build()
    assert ws.title == "2026年8月度 売上"


def test_title_row_text_and_merge():
    wb, ws = _build()
    assert ws["A1"].value == "㈱ケイピーエス　大阪支社　　2026年 8月度 売上表"
    assert any(str(r) == "A1:O1" for r in ws.merged_cells.ranges)


def test_header_rows_match_the_template(tpl):
    wb, ws = _build()
    for coord in ("A2", "B2", "C2", "E2", "G2", "I2", "J2", "K2", "L2",
                  "M2", "N2", "O2", "P2", "C3", "D3", "E3", "F3", "G3", "H3",
                  "O3", "P3"):
        assert ws[coord].value == tpl.ws[coord].value, f"{coord} が見出しと違う"


def test_header_merges_match_the_template():
    wb, ws = _build()
    ranges = {str(r) for r in ws.merged_cells.ranges}
    for rng in S.HEADER_MERGES:
        assert rng in ranges, f"{rng} が結合されていない"


# ----- 関西ぱど（京阪南・京阪北） -----


def test_pado_rows_are_fixed_at_rows_4_and_5():
    wb, ws = _build()
    assert ws["B4"].value == "京阪南"
    assert ws["B5"].value == "京阪北"


def test_pado_row_values_go_into_the_template_columns():
    wb, ws = _build()
    assert ws["J4"].value == 770456      # 売上合計(税抜) = 847502/1.1
    assert ws["O4"].value == 324505      # 配布原価(税込)
    assert ws["K4"].value == "=ROUND(J4*1.1,0)"   # テンプレートの数式をそのまま使う
    assert ws["P4"].value == "=M4+N4+O4"


# ----- 単独チラシ（「その他」区分・複数件あれば行を増やす） -----


def test_single_flyer_section_keeps_the_template_hint_when_there_is_no_data():
    wb, ws = _build()
    assert _find_row(ws, 1, "単独チラシ(あれば入力、案件ごとに行増やす)") is not None


def test_single_flyer_rows_grow_when_there_are_several_cases():
    """🔴 依頼⑦の例外対応: 単独案件が複数あれば行を増やす。"""
    wb, ws = _build(rows=[
        ("その他　買取専科", _row(176000, 132000)),
        ("その他　アクバススポーツクラブ", _row(343200, 200000)),
        ("その他　進和プロモーション", _row(112420, 57057)),
    ])
    for name in ("その他　買取専科", "その他　アクバススポーツクラブ",
                 "その他　進和プロモーション"):
        assert _find_row(ws, 1, name) is not None, f"{name} の行が無い"


def test_added_single_flyer_rows_do_not_break_the_summary_formula():
    """🔴 行を増やすと数式範囲がズレる。増やした行まで含めて集計されること。"""
    wb, ws = _build(rows=[
        ("関西ぱど：京阪南版", _row(1100, 0)),
        ("その他　買取専科", _row(2200, 0)),
        ("その他　配夢", _row(3300, 0)),
    ])
    r_minami = _find_row(ws, 2, "京阪南")
    r_a = _find_row(ws, 1, "その他　買取専科")
    r_b = _find_row(ws, 1, "その他　配夢")
    r_sum = _find_row(ws, 2, "北大阪営業部")
    formula = ws.cell(r_sum, 10).value          # J列(売上合計 税抜)の集計
    for r in (r_minami, r_a, r_b):
        assert f"J{r}" in formula, f"J{r} が集計の数式 {formula} に入っていない"


# ----- アドバリュー（週ごと） -----


def test_advalue_section_has_five_week_slots_labelled_for_that_month():
    """テンプレートと同じく 8-1〜8-5 の枠を出す。月をまたげば 9-1〜9-5 になる。"""
    wb, ws = _build()
    for w in range(1, 6):
        assert _find_row(ws, 1, f"ｱﾄﾞ・バリュー　8-{w}") is not None, f"8-{w} の枠が無い"


def test_advalue_week_slots_follow_the_month():
    wb, ws = _build(rows=[], year=2026, month=10)
    assert _find_row(ws, 1, "ｱﾄﾞ・バリュー　10-1") is not None
    assert _find_row(ws, 1, "ｱﾄﾞ・バリュー　10-5") is not None


def test_advalue_values_land_on_the_matching_week_row():
    wb, ws = _build()
    r1 = _find_row(ws, 1, "ｱﾄﾞ・バリュー　8-1")
    r2 = _find_row(ws, 1, "ｱﾄﾞ・バリュー　8-2")
    assert ws.cell(r1, 13).value == 6260        # M列 交通費
    assert ws.cell(r1, 15).value == 171280      # O列 配布原価(税込)=177540-6260
    assert ws.cell(r2, 13).value == 400
    assert ws.cell(r2, 10).value == round(214466 / 1.1)


def test_unused_advalue_week_rows_stay_blank():
    wb, ws = _build()
    r5 = _find_row(ws, 1, "ｱﾄﾞ・バリュー　8-5")
    assert ws.cell(r5, 10).value is None       # J 売上
    assert ws.cell(r5, 15).value is None       # O 配布原価


def test_extra_advalue_weeks_beyond_five_still_get_a_row():
    """6週目が来る月もある。枠から溢れた分を黙って捨てない。"""
    wb, ws = _build(rows=[("アドバリュー　8-6", _row(11000, 5000))])
    assert _find_row(ws, 1, "ｱﾄﾞ・バリュー　8-6") is not None


def test_advalue_labels_from_another_month_are_not_dropped():
    """月をまたいだ登録ミス等で 7-1 が混ざっても、行として残す。"""
    wb, ws = _build(rows=[("アドバリュー　7-1", _row(11000, 5000))])
    assert _find_row(ws, 1, "ｱﾄﾞ・バリュー　7-1") is not None


# ----- リビング -----


def test_living_section_has_the_two_rows_of_the_template():
    wb, ws = _build()
    r = _find_row(ws, 2, "リビング")
    assert r is not None
    assert any(str(rng) == f"A{r}:A{r + 1}" for rng in ws.merged_cells.ranges)
    assert any(str(rng) == f"B{r}:B{r + 1}" for rng in ws.merged_cells.ranges)


# ----- 未分類（案件マスタに無い/案件未設定） -----


def test_unclassified_rows_get_their_own_block_and_summary_line():
    """🔴 取りこぼしゼロ。どのセクションにも当てはまらない案件を捨てない。"""
    wb, ws = _build(rows=[("知らない案件", _row(5500, 1000))])
    r = _find_row(ws, 1, "知らない案件")
    assert r is not None
    r_sum = _find_row(ws, 2, "その他")
    assert r_sum is not None
    assert f"J{r}" in ws.cell(r_sum, 10).value


def test_no_unclassified_block_when_there_is_nothing_to_show():
    wb, ws = _build()
    assert _find_row(ws, 2, "その他") is None


# ----- 集計ブロック -----


def test_summary_block_has_the_four_lines_of_the_template():
    wb, ws = _build()
    labels = [ws.cell(r, 12).value for r in range(1, ws.max_row + 1)]
    for name in (" 関西ぱど", " アド・バリュー", " リビング・プロシード"):
        assert name in labels, f"集計の {name} 行が無い"
    assert _find_row(ws, 2, "計") is not None


def test_summary_block_month_label_is_merged_down_the_a_column():
    wb, ws = _build()
    r = _find_row(ws, 1, "8月度")
    assert r is not None
    last = _find_row(ws, 2, "計")
    assert any(str(rng) == f"A{r}:A{last}" for rng in ws.merged_cells.ranges)


def test_advalue_summary_sums_only_the_advalue_rows():
    wb, ws = _build()
    r_sum = next(r for r in range(1, ws.max_row + 1)
                 if ws.cell(r, 12).value == " アド・バリュー")
    formula = ws.cell(r_sum, 10).value
    r1 = _find_row(ws, 1, "ｱﾄﾞ・バリュー　8-1")
    r5 = _find_row(ws, 1, "ｱﾄﾞ・バリュー　8-5")
    assert f"J{r1}" in formula and f"J{r5}" in formula
    assert f"J{_find_row(ws, 2, '京阪南')}" not in formula


def test_grand_total_row_sums_the_cost_columns_over_every_data_row():
    wb, ws = _build()
    r_total = _find_row(ws, 2, "計")
    for col, letter in ((13, "M"), (14, "N"), (15, "O")):
        f = ws.cell(r_total, col).value
        assert f and f.startswith(f"=SUM({letter}"), f"{letter}列の総計が数式でない: {f!r}"


def test_summary_rows_use_the_template_tax_formula():
    wb, ws = _build()
    r_sum = _find_row(ws, 2, "北大阪営業部")
    assert ws.cell(r_sum, 11).value == f"=ROUND(J{r_sum}*1.1,0)"


# ----- 書式 -----


def test_column_widths_are_copied_from_template(tpl):
    wb, ws = _build()
    for col in ("A", "B", "L"):
        assert ws.column_dimensions[col].width == pytest.approx(
            tpl.ws.column_dimensions[col].width, abs=0.05)


def test_data_row_font_and_number_format_match_the_template(tpl):
    wb, ws = _build()
    src = tpl.ws.cell(S.STYLE_ROWS["pado"], 11)   # K列(売上合計 税込)の見本
    assert ws["K4"].font.name == src.font.name
    assert ws["K4"].font.sz == src.font.sz
    assert ws["K4"].number_format == src.number_format


def test_print_settings_are_landscape_and_fit_to_page():
    wb, ws = _build()
    assert ws.page_setup.orientation == "landscape"
    assert ws.print_area


def test_monthly_filename():
    assert S.monthly_filename(2026, 8) == "2026年8月度_売上表.xlsx"


def test_build_monthly_workbook_with_no_rows_still_has_the_whole_layout():
    """データが1件も無くても、テンプレートどおりの白紙が出ること
    (手で書き込む土台として使えるように)。"""
    wb, ws = _build(rows=[])
    assert ws["A2"].value == "発行号"
    assert ws["B4"].value == "京阪南"
    assert _find_row(ws, 1, "ｱﾄﾞ・バリュー　8-1") is not None
    assert _find_row(ws, 2, "計") is not None
